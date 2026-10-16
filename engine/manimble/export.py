"""Export the deck to PowerPoint (.pptx) and Keynote (.key).

PowerPoint: one slide per build (or per slide). In *editable* mode the text stays real, editable text boxes in the
original font, size and colour; everything else (backdrop, formulas, images, shapes, lines) is placed as pictures at
their exact positions. In *pictures* mode each page is the exact rendered frame.

Keynote: the .pptx is opened in Keynote through AppleScript and saved as a .key document."""

from __future__ import annotations

import os
import subprocess
import tempfile
from pathlib import Path

PT_PER_FONT_SIZE = 0.945      # manim font_size -> PowerPoint points on a 13.33in slide (calibrated on Helvetica Neue)


def ensure_keynote_running() -> None:
    """Start Keynote through LaunchServices (Apple Events alone cannot launch it from a background process)."""
    import time
    if subprocess.run(["pgrep", "-x", "Keynote"], capture_output=True).returncode == 0:
        return
    subprocess.run(["open", "-g", "-a", "Keynote"], capture_output=True)
    for _ in range(60):
        if subprocess.run(["pgrep", "-x", "Keynote"], capture_output=True).returncode == 0:
            break
        time.sleep(0.5)
    time.sleep(2.5)


def _hex(c: str | None):
    from pptx.dml.color import RGBColor
    c = (c or "#FFFFFF").lstrip("#")
    if len(c) != 6:
        c = "FFFFFF"
    return RGBColor(int(c[0:2], 16), int(c[2:4], 16), int(c[4:6], 16))


def build_pptx(app, path: str, builds: bool = False, editable: bool = True, progress=None) -> dict:
    try:
        from pptx import Presentation
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.util import Emu, Pt
    except ImportError:
        return {"ok": False, "error": "python-pptx is not installed in the Python environment Manimble uses "
                                      "(pip install python-pptx)."}
    prs = Presentation()
    blank = prs.slide_layouts[6]
    scenes = app.visible_scenes()
    total = len(scenes)
    first = True
    for si, s in enumerate(scenes):
        if progress:
            progress(si, total, s["name"])
        sc = app.api_scene(s["name"])
        if sc.get("status") != "ready":
            return {"ok": False, "error": f"{s['name']} does not render: {str(sc.get('error', ''))[-300:]}"}
        pw, ph = sc["size"]
        if first:
            prs.slide_width, prs.slide_height = Emu(12192000), Emu(int(12192000 * ph / pw))
            first = False
        k = 12192000 / pw                                   # EMU per frame pixel
        stops = sc["stops"] if builds else sc["stops"][-1:]
        for st in stops:
            slide = prs.slides.add_slide(blank)
            if not editable:
                slide.shapes.add_picture(st["png"], 0, 0, prs.slide_width, prs.slide_height)
                continue
            slide.shapes.add_picture(st["bg"], 0, 0, prs.slide_width, prs.slide_height)
            vis = {int(v[0]): v for v in st["visible"]}
            for sp in st["sprites"]:
                oid = sp["oid"]
                obj = sc["objects"].get(str(oid), {})
                v = vis.get(oid)
                if obj.get("kind") == "text" and v is not None and obj.get("value"):
                    x0, y0, x1, y1 = v[1:5]
                    padx = 0.06 * (x1 - x0) + 30
                    box = slide.shapes.add_textbox(Emu(int((x0 - padx) * k)), Emu(int((y0 - 8) * k)),
                                                   Emu(int((x1 - x0 + 2 * padx) * k)), Emu(int((y1 - y0 + 16) * k)))
                    tf = box.text_frame
                    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                    tf.word_wrap = False
                    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                    lines = str(obj["value"]).split("\n")
                    for li, line in enumerate(lines):
                        para = tf.paragraphs[0] if li == 0 else tf.add_paragraph()
                        para.alignment = PP_ALIGN.CENTER
                        run = para.add_run()
                        run.text = line
                        f = run.font
                        if obj.get("font"):
                            f.name = obj["font"]
                        f.size = Pt(max(4, round((obj.get("fontSize") or 30) * PT_PER_FONT_SIZE, 1)))
                        f.bold = str(obj.get("weight") or "").upper() == "BOLD"
                        f.italic = str(obj.get("slant") or "").upper() in ("ITALIC", "OBLIQUE")
                        f.color.rgb = _hex(obj.get("color"))
                else:
                    slide.shapes.add_picture(sp["file"], Emu(int(sp["x"] * k)), Emu(int(sp["y"] * k)),
                                             Emu(int(sp["w"] * k)), Emu(int(sp["h"] * k)))
    prs.save(path)
    return {"ok": True, "path": path, "slides": len(prs.slides)}


def pptx_to_keynote(pptx_path: str, key_path: str) -> dict:
    """Open the .pptx in Keynote and save it as a Keynote document."""
    if not Path("/Applications/Keynote.app").exists():
        return {"ok": False, "error": "Keynote is not installed; open the exported .pptx in Keynote instead."}
    ensure_keynote_running()
    script = f'''
    set src to POSIX file "{pptx_path}"
    set dst to POSIX file "{key_path}"
    tell application "Keynote"
        activate
        set d to open src
        delay 2
        save d in dst
        close d saving no
    end tell
    '''
    try:
        r = subprocess.run(["osascript", "-e", script], capture_output=True, text=True, timeout=240)
    except subprocess.TimeoutExpired:
        return {"ok": False, "error": "Keynote did not respond (a dialog may be waiting in Keynote)."}
    if r.returncode != 0 or not Path(key_path).exists():
        return {"ok": False, "error": "Keynote could not save the file: " + (r.stderr or r.stdout).strip()[-300:]}
    return {"ok": True, "path": key_path}
