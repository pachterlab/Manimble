"""Import a PowerPoint (.pptx/.ppt) or Keynote (.key) presentation as a new Manim Slides deck.

Every slide becomes a scene class. Text boxes become `Text(...)`, pictures `ImageMobject(...)`, rectangles /
ellipses / lines / block arrows the matching manim shapes, tables `Table(...)`. Elements are written one per line,
directly in the scene, so Manimble can edit them like anything else. Animations and transitions are not imported.

Positions, sizes and styles are resolved the way PowerPoint resolves them: group transforms (offset, scale,
rotation, flips) are applied to the shapes inside a group, theme colours and the colour map are looked up, shape
styles (p:style) supply fills, outlines and text colours that the shape does not set itself, and text properties
that a run does not set are inherited from the shape's list style, its placeholder and the master's text styles."""

from __future__ import annotations

import colorsys
import glob
import io
import math
import re
import subprocess
import tempfile
from pathlib import Path

FW, FH = 14.222222222222221, 8.0          # manim frame in units; 1280 x 720 px
DEFAULT_FONT = "Helvetica Neue"          # used when neither the presentation nor its theme names a font
PT_PER_FONT_SIZE = 0.945                  # same calibration as the exporter

# Fonts that ship inside Microsoft Office (Calibri, Cambria, ...) are not installed system-wide, but can be used
# for measuring text here and registered for rendering by the deck (FONT_REGISTER_SOURCE).
OFFICE_FONT_DIRS = tuple(sorted(glob.glob("/Applications/Microsoft *.app/Contents/Resources/DFonts")))
_FONT_DIRS = ("/System/Library/Fonts", "/System/Library/Fonts/Supplemental", "/Library/Fonts",
              str(Path.home() / "Library" / "Fonts")) + OFFICE_FONT_DIRS

# written into an imported deck's common.py when it has a table
TABLE_GRID_SOURCE = '''

def _baseline(line):
    bottoms = sorted(g.get_bottom()[1] for g in line.family_members_with_points())
    return bottoms[len(bottoms) // 2]


def table_grid(table, col_widths, row_heights, align=UL, margin=(0.1, 0.05)):
    """Lay a Table out the way PowerPoint does: fixed column widths and minimum row heights (in units),
    each entry aligned in its cell by `align` inside `margin` with the text on common baselines, the lines
    moved onto the cell edges."""
    cells = table.mob_table
    n_r, n_c = len(cells), len(cells[0])
    if len(col_widths) != n_c:          # columns added or removed since the import: share the width
        col_widths = [sum(col_widths) / n_c] * n_c
    if len(row_heights) != n_r:
        row_heights = [sum(row_heights) / len(row_heights)] * n_r
    mx, my = margin
    # each entry's text block runs from the tallest glyph's height above its first baseline down to its last
    # baseline, so cells with and without ascenders line up
    lines = {id(m): [l for l in m.submobjects if l.family_members_with_points()] if isinstance(m, Paragraph) else []
             for row in cells for m in row}
    asc = max([l[0].get_top()[1] - _baseline(l[0]) for l in lines.values() if l] or [0])

    def block(m):                        # (top, bottom) of the text block
        ls = lines[id(m)]
        return (_baseline(ls[0]) + asc, _baseline(ls[-1])) if ls else (m.get_top()[1], m.get_bottom()[1])
    row_heights = [max(h, max(t - b for t, b in map(block, row)) + 2 * my) for h, row in zip(row_heights, cells)]
    xs, ys = [0.0], [0.0]
    for w in col_widths:
        xs.append(xs[-1] + w)
    for h in row_heights:
        ys.append(ys[-1] - h)
    for r, row in enumerate(cells):
        for c, m in enumerate(row):
            x = xs[c] + mx if align[0] < 0 else xs[c + 1] - mx if align[0] > 0 else (xs[c] + xs[c + 1]) / 2
            m.move_to((x, 0, 0), np.array([align[0], 0, 0]))
            t, b = block(m)
            top = ys[r] - my if align[1] > 0 else ys[r + 1] + my + t - b if align[1] < 0 else (ys[r] + ys[r + 1] + t - b) / 2
            m.shift((top - t) * UP)
    # move the Table's own lines (outer ones first, as Table makes them) rather than drawing new ones
    order = (lambda v: [v[0], v[-1], *v[1:-1]]) if table.include_outer_lines else (lambda v: v[1:-1])
    for line, x in zip(table.vertical_lines, order(xs)):
        line.put_start_and_end_on((x, ys[0], 0), (x, ys[-1], 0))
    for line, y in zip(table.horizontal_lines, order(ys)):
        line.put_start_and_end_on((xs[0], y, 0), (xs[-1], y, 0))
    return table
'''

# written into an imported deck's common.py: fonts the presentation uses that are not installed system-wide but
# ship with Microsoft Office are registered for the rendering process, so manim does not substitute another font
FONT_REGISTER_SOURCE = '''

def _register_fonts(*names):
    """Make fonts that are not installed, but ship with Microsoft Office, available to manim in this process."""
    try:
        import glob
        import manimpango
        have = {f.lower() for f in manimpango.list_fonts()}
    except Exception:
        return
    for name in names:
        if name.lower() in have:
            continue
        stem = name.split()[0].lower()
        for d in glob.glob("/Applications/Microsoft *.app/Contents/Resources/DFonts"):
            for f in Path(d).iterdir():
                if f.suffix.lower() in (".ttf", ".otf", ".ttc") and f.stem.lower().startswith(stem):
                    try:
                        manimpango.register_font(str(f))
                    except Exception:
                        pass
'''


def keynote_to_pptx(key_path: str, pptx_path: str) -> dict:
    from .export import ensure_keynote_running
    if not Path("/Applications/Keynote.app").exists():
        return {"ok": False, "error": "Keynote is not installed, so a .key or .ppt file cannot be read. "
                                      "Save it as .pptx and import that instead."}
    ensure_keynote_running()
    script = f'''
    set src to POSIX file "{key_path}"
    set dst to POSIX file "{pptx_path}"
    tell application "Keynote"
        activate
        set d to open src
        delay 2
        export d to dst as Microsoft PowerPoint
        close d saving no
    end tell
    '''
    try:
        r = subprocess.run(["osascript", "-e", script], capture_output=True, text=True, timeout=300)
    except subprocess.TimeoutExpired:
        return {"ok": False, "error": "Keynote did not respond (a dialog may be waiting in Keynote)."}
    if r.returncode != 0 or not Path(pptx_path).exists():
        return {"ok": False, "error": "Keynote could not open the file: " + (r.stderr or r.stdout).strip()[-300:]}
    return {"ok": True}


# ----------------------------------------------------------------------------------------------- small helpers

def _luma(hexcolor: str) -> float:
    h = hexcolor.lstrip("#")
    r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    return 0.299 * r + 0.587 * g + 0.114 * b


def _slug(s: str, fallback: str) -> str:
    s = re.sub(r"[^A-Za-z0-9]+", "_", s).strip("_")[:26]
    return s or fallback


def _num(x: float, nd: int = 3) -> str:
    s = f"{x:.{nd}f}".rstrip("0").rstrip(".")
    return "0" if s in ("", "-0") else s


def _local(el) -> str:
    return el.tag.rsplit("}", 1)[-1]


def _int(v, default=0) -> int:
    try:
        return int(v)
    except (TypeError, ValueError):
        return default


_font_cache: dict = {}


def _font(family: str, bold: bool, italic: bool):
    """PIL font (at 1000 px per em) for measuring text the way PowerPoint wraps it, or None if not installed."""
    key = (family.lower(), bold, italic)
    if key in _font_cache:
        return _font_cache[key]
    want = ("bold " if bold else "") + ("italic" if italic else "")
    want = want.strip() or "regular"
    stem = family.split()[0].lower() if family.split() else family.lower()
    found = None
    try:
        from PIL import ImageFont
        faces = []
        for d in _FONT_DIRS:
            p = Path(d)
            if not p.is_dir():
                continue
            for f in p.iterdir():
                if f.suffix.lower() in (".ttf", ".ttc", ".otf") and f.stem.lower().replace(" ", "").startswith(stem):
                    for i in range(64):
                        try:
                            ft = ImageFont.truetype(str(f), 1000, index=i)
                        except Exception:                       # noqa: BLE001
                            break
                        fam, style = ft.getname()
                        if (fam or "").lower() == family.lower():
                            faces.append(((style or "").lower().replace("oblique", "italic"), ft))
        found = next((ft for style, ft in faces if style == want), None)
        if found is None and faces:
            found = next((ft for style, ft in faces if style == "regular"), faces[0][1])
    except Exception:                                           # noqa: BLE001
        found = None
    _font_cache[key] = found
    return found


def _font_or_fallback(family: str, bold: bool, italic: bool):
    """The font for measuring, or a common sans-serif when the family is not available anywhere (what Pango
    will substitute is unknown, but Arial's widths are closer than a fixed average)."""
    return _font(family, bold, italic) or _font("Arial", bold, italic) or _font("Helvetica", bold, italic)


def _wrap(text: str, styles: list, width_em: float, measure) -> tuple[str, list, list]:
    """Greedy word wrap to width_em (in ems of the font size), like PowerPoint: break at spaces and after
    hyphens, and inside a word only when the word alone is wider than the box. styles[i] is the
    style of text[i]. Returns (text, styles, idx): idx[k] is the index in the original text of result[k], or -1
    for an inserted line break. measure(segment, style) -> ems."""
    def width(i, j):
        w, k = 0.0, i
        while k < j:
            m = k
            while m < j and styles[m] == styles[k]:
                m += 1
            w += measure(text[k:m], styles[k])
            k = m
        return w

    tokens, start = [], 0
    for i, ch in enumerate(text):
        if ch == " ":
            tokens.append((start, i + 1, "space"))
            start = i + 1
        elif ch == "\n":
            tokens.append((start, i + 1, "hard"))
            start = i + 1
        elif ch == "-" and i + 1 < len(text) and text[i + 1] not in " \n":
            tokens.append((start, i + 1, "soft"))
            start = i + 1
    if start < len(text):
        tokens.append((start, len(text), "end"))
    out: list = []                                  # (char, style, original index)
    line_w = 0.0
    for s, e, kind in tokens:
        body_end = e - 1 if kind in ("space", "hard") else e
        tw = width(s, body_end)
        if line_w > 0 and line_w + tw > width_em + 1e-9:
            if out and out[-1][0] == " ":
                out[-1] = ("\n", out[-1][1], out[-1][2])
            elif out and out[-1][0] != "\n":
                out.append(("\n", styles[s] if s < len(styles) else None, -1))
            line_w = 0.0
        if tw > width_em:                           # wider than the box on its own: break between characters
            for k in range(s, body_end):
                cw = measure(text[k], styles[k])
                if line_w > 0 and line_w + cw > width_em:
                    out.append(("\n", styles[k], -1))
                    line_w = 0.0
                out.append((text[k], styles[k], k))
                line_w += cw
        else:
            out.extend((text[k], styles[k], k) for k in range(s, body_end))
            line_w += tw
        if kind == "space":
            out.append((" ", styles[e - 1], e - 1))
            line_w += measure(" ", styles[e - 1])
        elif kind == "hard":
            out.append(("\n", styles[e - 1], e - 1))
            line_w = 0.0
    return "".join(c for c, _, _ in out), [st for _, st, _ in out], [i for _, _, i in out]


# ------------------------------------------------------------------------------------- Office Math (OMML)

_MATH_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
_GREEK = "ΑΒΓΔΕΖΗΘΙΚΛΜΝΞΟΠΡϴΣΤΥΦΧΨΩ∇αβγδεζηθικλμνξοπρςστυφχψω∂ϵϑϰϕϱϖ"
_LATEX_SYM = {
    "∝": r"\propto ", "λ": r"\lambda ", "μ": r"\mu ", "σ": r"\sigma ", "α": r"\alpha ", "β": r"\beta ",
    "γ": r"\gamma ", "δ": r"\delta ", "ε": r"\varepsilon ", "ϵ": r"\epsilon ", "θ": r"\theta ", "π": r"\pi ",
    "ρ": r"\rho ", "τ": r"\tau ", "φ": r"\varphi ", "ϕ": r"\phi ", "ω": r"\omega ", "ψ": r"\psi ", "χ": r"\chi ",
    "η": r"\eta ", "κ": r"\kappa ", "ν": r"\nu ", "ξ": r"\xi ", "ζ": r"\zeta ", "ι": r"\iota ", "υ": r"\upsilon ",
    "ς": r"\varsigma ", "ϑ": r"\vartheta ", "ϰ": r"\varkappa ", "ϱ": r"\varrho ", "ϖ": r"\varpi ",
    "Σ": r"\Sigma ", "Δ": r"\Delta ", "Ω": r"\Omega ", "Γ": r"\Gamma ", "Θ": r"\Theta ", "Λ": r"\Lambda ",
    "Π": r"\Pi ", "Φ": r"\Phi ", "Ψ": r"\Psi ", "Ξ": r"\Xi ", "Υ": r"\Upsilon ", "∇": r"\nabla ", "∂": r"\partial ",
    "∗": "*", "×": r"\times ", "·": r"\cdot ", "⋅": r"\cdot ", "≈": r"\approx ", "≠": r"\neq ", "≤": r"\le ",
    "≥": r"\ge ", "≪": r"\ll ", "≫": r"\gg ", "∞": r"\infty ", "→": r"\to ", "←": r"\leftarrow ",
    "⇒": r"\Rightarrow ", "⇔": r"\Leftrightarrow ", "∈": r"\in ", "∉": r"\notin ", "⊂": r"\subset ",
    "⊆": r"\subseteq ", "∪": r"\cup ", "∩": r"\cap ", "±": r"\pm ", "∓": r"\mp ", "∑": r"\sum ", "∏": r"\prod ",
    "∫": r"\int ", "…": r"\ldots ", "⋯": r"\cdots ", "′": "'", "°": r"^\circ ", "∼": r"\sim ", "∧": r"\wedge ",
    "∨": r"\vee ", "¬": r"\neg ", "∀": r"\forall ", "∃": r"\exists ", "∅": r"\emptyset ", "ℝ": r"\mathbb{R}",
    "ℕ": r"\mathbb{N}", "ℤ": r"\mathbb{Z}", "ℚ": r"\mathbb{Q}", "ℂ": r"\mathbb{C}", "ℓ": r"\ell ",
    "−": "-", "%": r"\%", "&": r"\&", "#": r"\#", "$": r"\$", "{": r"\{", "}": r"\}", "~": r"\sim ",
    "\\": r"\backslash ", "\u2009": r"\,", "\u2061": "", "\u2062": "", "\u2063": "", "\xa0": r"\ "}
_SUP = {"0": "⁰", "1": "¹", "2": "²", "3": "³", "4": "⁴", "5": "⁵", "6": "⁶", "7": "⁷", "8": "⁸", "9": "⁹",
        "+": "⁺", "-": "⁻", "−": "⁻", "=": "⁼", "(": "⁽", ")": "⁾", "n": "ⁿ", "i": "ⁱ", "T": "ᵀ", "t": "ᵗ",
        "x": "ˣ", "k": "ᵏ", "m": "ᵐ", "a": "ᵃ", "b": "ᵇ", "c": "ᶜ", "d": "ᵈ", "e": "ᵉ", "j": "ʲ", "p": "ᵖ",
        "r": "ʳ", "s": "ˢ", "u": "ᵘ", "y": "ʸ", "N": "ᴺ", "M": "ᴹ", "K": "ᴷ", "L": "ᴸ", "H": "ᴴ"}
_SUB = {"0": "₀", "1": "₁", "2": "₂", "3": "₃", "4": "₄", "5": "₅", "6": "₆", "7": "₇", "8": "₈", "9": "₉",
        "+": "₊", "-": "₋", "−": "₋", "=": "₌", "(": "₍", ")": "₎", "a": "ₐ", "e": "ₑ", "h": "ₕ", "i": "ᵢ",
        "j": "ⱼ", "k": "ₖ", "l": "ₗ", "m": "ₘ", "n": "ₙ", "o": "ₒ", "p": "ₚ", "r": "ᵣ", "s": "ₛ", "t": "ₜ",
        "u": "ᵤ", "v": "ᵥ", "x": "ₓ"}
_DELIM = {"{": r"\{", "}": r"\}", "⟨": r"\langle ", "⟩": r"\rangle ", "‖": r"\|", "⌈": r"\lceil ",
          "⌉": r"\rceil ", "⌊": r"\lfloor ", "⌋": r"\rfloor ", "": "."}
_NARY = {"∑": r"\sum", "∏": r"\prod", "∐": r"\coprod", "∫": r"\int", "∬": r"\iint", "∭": r"\iiint",
         "∮": r"\oint", "⋃": r"\bigcup", "⋂": r"\bigcap", "⋁": r"\bigvee", "⋀": r"\bigwedge"}
_ACC = {"\u0302": "hat", "\u0303": "tilde", "\u0304": "bar", "\u0305": "bar", "\u20d7": "vec", "\u0307": "dot",
        "\u0308": "ddot", "\u0306": "breve", "\u030c": "check", "\u0301": "acute", "\u0300": "grave"}
_FUNCS = {"sin", "cos", "tan", "log", "ln", "exp", "max", "min", "lim", "det", "arg", "sup", "inf", "sinh",
          "cosh", "tanh", "sec", "csc", "cot", "arcsin", "arccos", "arctan", "gcd", "deg", "dim", "ker", "Pr"}


def _demath(s: str) -> str:
    """Mathematical alphanumeric symbols (the italic/bold letters Office Math uses) as plain letters."""
    out = []
    for ch in s:
        cp = ord(ch)
        if 0x1D400 <= cp <= 0x1D6A3:
            i = (cp - 0x1D400) % 52
            out.append(chr(ord("A") + i) if i < 26 else chr(ord("a") + i - 26))
        elif 0x1D6A8 <= cp <= 0x1D7C9:
            out.append(_GREEK[(cp - 0x1D6A8) % 58])
        elif 0x1D7CE <= cp <= 0x1D7FF:
            out.append(str((cp - 0x1D7CE) % 10))
        elif ch == "ℎ":
            out.append("h")
        else:
            out.append(ch)
    return "".join(out)


def _script(s: str, table: dict, mark: str) -> str:
    if not s:
        return ""
    if all(ch in table for ch in s):
        return "".join(table[ch] for ch in s)
    return mark + (s if len(s) == 1 else "(" + s + ")")


def _paren(s: str) -> str:
    return s if len(s) <= 1 else "(" + s + ")"


def _om(el, latex: bool) -> str:
    """An Office Math element as LaTeX (latex=True) or as plain text that reads about right in a Text."""
    M = _MATH_NS
    t = _local(el)

    def part(tag):
        c = el.find(M + tag)
        return _om(c, latex) if c is not None else ""

    def val(e, tag, default=None):
        if e is None:
            return default
        c = e.find(M + tag)
        return c.get(M + "val", default) if c is not None else default

    pr = el.find(M + t + "Pr")
    if t == "r":
        txt = _demath("".join(x.text or "" for x in el.iter(M + "t")))
        rpr = el.find(M + "rPr")
        upright = rpr is not None and (rpr.find(M + "nor") is not None or val(rpr, "sty") in ("p", "b"))
        if not latex:
            return txt
        s = "".join(r"\ " if ch == " " and upright else _LATEX_SYM.get(ch, ch) for ch in txt)
        return r"\mathrm{" + s + "}" if upright else s
    if t in ("sSup", "sSub", "sSubSup", "sPre"):
        e, sub, sup = part("e"), part("sub"), part("sup")
        if latex:
            scripts = (f"_{{{sub}}}" if sub else "") + (f"^{{{sup}}}" if sup else "")
            return "{}" + scripts + f"{{{e}}}" if t == "sPre" else f"{{{e}}}" + scripts
        return e + _script(sub, _SUB, "_") + _script(sup, _SUP, "^")
    if t == "f":
        num, den = part("num"), part("den")
        if latex:
            return f"{num}/{den}" if val(pr, "type") in ("lin", "skw") else rf"\frac{{{num}}}{{{den}}}"
        return f"{_paren(num)}/{_paren(den)}"
    if t == "d":
        beg, end, sep = val(pr, "begChr", "("), val(pr, "endChr", ")"), val(pr, "sepChr", "|")
        es = [_om(c, latex) for c in el.findall(M + "e")]
        if latex:
            return r"\left" + _DELIM.get(beg, beg) + sep.join(es) + r"\right" + _DELIM.get(end, end)
        return beg + sep.join(es) + end
    if t == "nary":
        ch = val(pr, "chr", "∫")
        sub = "" if val(pr, "subHide") == "on" else part("sub")
        sup = "" if val(pr, "supHide") == "on" else part("sup")
        e = part("e")
        if latex:
            return _NARY.get(ch, r"\int") + (f"_{{{sub}}}" if sub else "") + (f"^{{{sup}}}" if sup else "") + f"{{{e}}}"
        return ch + _script(sub, _SUB, "_") + _script(sup, _SUP, "^") + _paren(e)
    if t == "rad":
        deg = "" if val(pr, "degHide") == "on" else part("deg")
        e = part("e")
        if latex:
            return rf"\sqrt[{deg}]{{{e}}}" if deg else rf"\sqrt{{{e}}}"
        return "√" + _paren(e)
    if t == "acc":
        ch, e = val(pr, "chr", "\u0302"), part("e")
        return "\\" + _ACC.get(ch, "hat") + f"{{{e}}}" if latex else e + ch
    if t == "bar":
        e = part("e")
        if latex:
            return (r"\overline{" if val(pr, "pos") == "top" else r"\underline{") + e + "}"
        return e
    if t == "func":
        name, e = part("fName"), part("e")
        if latex:
            plain = re.sub(r"\\mathrm\{([A-Za-z]+)\}", r"\1", name).strip()
            if plain in _FUNCS:
                name = "\\" + plain + " "
            return name + f"{{{e}}}"
        return name + e
    if t in ("limLow", "limUpp"):
        e, lim = part("e"), part("lim")
        if latex:
            return (r"\underset{" if t == "limLow" else r"\overset{") + lim + "}{" + e + "}"
        return e + "_" + lim
    if t == "groupChr":
        e = part("e")
        return r"\underbrace{" + e + "}" if latex else e
    if t == "m":
        rows = [" & ".join(_om(c, latex) for c in mr.findall(M + "e")) for mr in el.findall(M + "mr")]
        if latex:
            return r"\begin{matrix}" + r" \\ ".join(rows) + r"\end{matrix}"
        return "; ".join(r.replace(" & ", ", ") for r in rows)
    if t == "eqArr":
        es = [_om(c, latex) for c in el.findall(M + "e")]
        return (r" \\ " if latex else "; ").join(es)
    if t.endswith("Pr"):
        return ""
    return "".join(_om(c, latex) for c in el)


def omml_latex(el) -> str:
    s = _om(el, True).strip()
    return s + " " if s.endswith("\\") else s


def omml_text(el) -> str:
    return _om(el, False)


def omml_tall(el) -> bool:
    """Whether the formula has parts that stack (fractions, big operators, roots)."""
    return any(_local(c) in ("f", "nary", "rad", "m", "eqArr", "limLow", "limUpp") for c in el.iter())


# --------------------------------------------------------------------------------------------- theme & colours

_PRESET_COLORS = {"black": "000000", "white": "FFFFFF", "red": "FF0000", "green": "008000", "blue": "0000FF",
                  "yellow": "FFFF00", "gray": "808080", "grey": "808080", "lime": "00FF00", "navy": "000080",
                  "orange": "FFA500", "purple": "800080", "silver": "C0C0C0", "teal": "008080", "maroon": "800000",
                  "olive": "808000", "aqua": "00FFFF", "cyan": "00FFFF", "fuchsia": "FF00FF", "magenta": "FF00FF"}
_CLR_MAP_DEFAULT = {"tx1": "dk1", "bg1": "lt1", "tx2": "dk2", "bg2": "lt2"}
_CLR_TAGS = ("srgbClr", "schemeClr", "sysClr", "prstClr", "hslClr", "scrgbClr")
_FILL_TAGS = ("noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill")
_BULLETS = {("Wingdings", "§"): "▪", ("Wingdings", "q"): "❑", ("Wingdings", "Ø"): "➢", ("Wingdings", "ü"): "✓",
            ("Wingdings", "v"): "❖", ("Wingdings", "l"): "●", ("Wingdings", "n"): "■", ("Wingdings", "u"): "◆",
            ("Wingdings", "Ü"): "➔", ("Wingdings", "ñ"): "◆", ("Wingdings", "¨"): "▫", ("Symbol", "·"): "•",
            ("Symbol", "-"): "–", ("Courier New", "o"): "○", ("Wingdings 2", "Ú"): "✖"}


def _lin(c: float) -> float:
    return c ** 2.2


def _lin_inv(c: float) -> float:
    return max(0.0, min(1.0, c)) ** (1 / 2.2)


class _Theme:
    """Colours, fonts and default line widths of a slide master's theme, plus its text styles."""

    def __init__(self, master, qn):
        self.qn = qn
        self.colors: dict = {}
        self.fonts = {"major": "Calibri Light", "minor": "Calibri"}
        self.line_widths = [9525, 9525, 19050, 28575]
        self.clr_map: dict = dict(_CLR_MAP_DEFAULT)
        self.title_style = self.body_style = self.other_style = None
        try:
            from lxml import etree
            from pptx.opc.constants import RELATIONSHIP_TYPE as RT
            root = etree.fromstring(master.part.part_related_by(RT.THEME).blob)
            cs = root.find(".//" + qn("a:clrScheme"))
            if cs is not None:
                for child in cs:
                    if len(child):
                        c = child[0]
                        self.colors[_local(child)] = (c.get("val") if _local(c) == "srgbClr" else c.get("lastClr")) or "000000"
            for kind in ("major", "minor"):
                lat = root.find(f".//{qn('a:' + kind + 'Font')}/{qn('a:latin')}")
                if lat is not None and lat.get("typeface"):
                    self.fonts[kind] = lat.get("typeface")
            lst = root.find(".//" + qn("a:lnStyleLst"))
            if lst is not None and len(lst):
                self.line_widths = [9525] + [_int(ln.get("w"), 9525) for ln in lst]
        except Exception:                                       # noqa: BLE001
            pass
        cm = master._element.find(qn("p:clrMap"))
        if cm is not None:
            self.clr_map.update(cm.attrib)
        tx = master._element.find(qn("p:txStyles"))
        if tx is not None:
            self.title_style = tx.find(qn("p:titleStyle"))
            self.body_style = tx.find(qn("p:bodyStyle"))
            self.other_style = tx.find(qn("p:otherStyle"))

    def font(self, typeface: str | None) -> str | None:
        if not typeface:
            return None
        if typeface.startswith("+mj"):
            return self.fonts["major"]
        if typeface.startswith("+mn"):
            return self.fonts["minor"]
        return typeface

    def clr(self, el, ph: str | None = None):
        """(hex, alpha) for a colour element (srgbClr / schemeClr / ...), with its modifiers applied."""
        tag = _local(el)
        if tag == "srgbClr":
            rgb = el.get("val")
        elif tag == "sysClr":
            rgb = el.get("lastClr") or "000000"
        elif tag == "schemeClr":
            name = el.get("val") or ""
            if name == "phClr":
                rgb = ph
            else:
                rgb = self.colors.get(self.clr_map.get(name, name)) or self.colors.get(name)
        elif tag == "prstClr":
            rgb = _PRESET_COLORS.get(el.get("val") or "", "000000")
        elif tag == "scrgbClr":
            try:
                rgb = "".join("%02X" % round(_lin_inv(_int(el.get(k)) / 100000) * 255) for k in ("r", "g", "b"))
            except Exception:                                   # noqa: BLE001
                rgb = None
        else:
            rgb = None
        if not rgb or len(rgb) < 6:
            return None
        try:
            r, g, b = (int(rgb[i:i + 2], 16) / 255 for i in (0, 2, 4))
        except ValueError:
            return None
        alpha = 1.0
        for m in el:
            t = _local(m)
            v = _int(m.get("val")) / 100000
            if t == "alpha":
                alpha = v
            elif t in ("lumMod", "lumOff", "satMod", "satOff"):
                h, l, s = colorsys.rgb_to_hls(r, g, b)
                if t == "lumMod":
                    l *= v
                elif t == "lumOff":
                    l += v
                elif t == "satMod":
                    s *= v
                else:
                    s += v
                r, g, b = colorsys.hls_to_rgb(h, min(max(l, 0.0), 1.0), min(max(s, 0.0), 1.0))
            elif t == "shade":
                r, g, b = (_lin_inv(_lin(c) * v) for c in (r, g, b))
            elif t == "tint":
                r, g, b = (_lin_inv(1 - (1 - _lin(c)) * v) for c in (r, g, b))
        return "#%02X%02X%02X" % (round(r * 255), round(g * 255), round(b * 255)), alpha

    def first_clr(self, parent, ph=None):
        if parent is None:
            return None
        for c in parent:
            if _local(c) in _CLR_TAGS:
                return self.clr(c, ph)
        return None

    def fill(self, tag: str, el, ph=None):
        """(hex, alpha) of a fill element, or None for no fill / a picture fill."""
        if tag == "solidFill":
            return self.first_clr(el, ph)
        if tag == "gradFill":
            stops = el.find(self.qn("a:gsLst"))
            if stops is not None and len(stops):
                return self.first_clr(stops[0], ph)
            return None
        if tag == "pattFill":
            return self.first_clr(el.find(self.qn("a:fgClr")), ph)
        return None


# ---------------------------------------------------------------------------------------------- group transform

class _Xf:
    """Maps coordinates in a group's child space onto the slide: each nesting level scales the children's
    coordinate space (chOff/chExt) onto the group's box (off/ext), then flips and rotates about its centre."""

    def __init__(self, parent=None, xfrm=None, qn=None):
        self.parent = parent
        self.scale_x = self.scale_y = 1.0
        self.rot = 0.0
        self.flip_h = self.flip_v = False
        self.ox = self.oy = self.cox = self.coy = self.cx = self.cy = 0.0
        self.kx = self.ky = 1.0
        if parent is not None:
            self.scale_x, self.scale_y = parent.scale_x, parent.scale_y
        if parent is not None and xfrm is not None:
            off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
            choff, chext = xfrm.find(qn("a:chOff")), xfrm.find(qn("a:chExt"))
            if off is not None and ext is not None:
                ox, oy = _int(off.get("x")), _int(off.get("y"))
                ex, ey = _int(ext.get("cx")), _int(ext.get("cy"))
                cox, coy = (_int(choff.get("x")), _int(choff.get("y"))) if choff is not None else (ox, oy)
                cex, cey = (_int(chext.get("cx")), _int(chext.get("cy"))) if chext is not None else (ex, ey)
                self.kx = ex / cex if cex else 1.0
                self.ky = ey / cey if cey else 1.0
                self.ox, self.oy, self.cox, self.coy = ox, oy, cox, coy
                self.cx, self.cy = ox + ex / 2, oy + ey / 2
                self.rot = _int(xfrm.get("rot")) / 60000
                self.flip_h, self.flip_v = xfrm.get("flipH") == "1", xfrm.get("flipV") == "1"
                self.scale_x = abs(self.kx) * parent.scale_x
                self.scale_y = abs(self.ky) * parent.scale_y

    def pt(self, x: float, y: float) -> tuple[float, float]:
        if self.parent is None:
            return x, y
        gx = self.ox + (x - self.cox) * self.kx
        gy = self.oy + (y - self.coy) * self.ky
        if self.flip_h:
            gx = 2 * self.cx - gx
        if self.flip_v:
            gy = 2 * self.cy - gy
        if self.rot:
            a = math.radians(self.rot)
            dx, dy = gx - self.cx, gy - self.cy
            gx = self.cx + dx * math.cos(a) - dy * math.sin(a)
            gy = self.cy + dx * math.sin(a) + dy * math.cos(a)
        return self.parent.pt(gx, gy)

    def angle(self, deg: float) -> float:
        """A shape's rotation as seen on the slide (degrees, clockwise), given its own rotation."""
        node, total = self, deg
        while node.parent is not None:
            if node.flip_h != node.flip_v:
                total = -total
            total += node.rot
            node = node.parent
        return total


class _Box:
    """A shape's placement on the slide: centre (EMU), size (EMU, after group scaling), rotation (degrees,
    clockwise) and flips; plus the raw xfrm values for lines."""

    def __init__(self, cx, cy, w, h, angle, flip_h, flip_v, raw):
        self.cx, self.cy, self.w, self.h, self.angle = cx, cy, w, h, angle
        self.flip_h, self.flip_v, self.raw = flip_h, flip_v, raw


# ------------------------------------------------------------------------------------------------- the importer

class _Importer:
    def __init__(self, prs, dest: Path, qn, enums):
        self.prs, self.qn = prs, qn
        self.MSO_SHAPE_TYPE, self.PP_ALIGN, self.MSO_ANCHOR = enums
        self.deck, self.assets = dest / "deck", dest / "assets"
        self.W, self.H = prs.slide_width, prs.slide_height
        self.s_ref = 12192000 / self.W                         # normalise to a 13.33in wide slide
        self.px_per_emu = 1280 / self.W
        self.warnings: list[str] = []
        self.n_img = 0
        self.uses_tables = False
        self.fonts_used: set = set()
        self.themes: dict = {}

    # ---------------------------------------------------------------------------------------- units & colours
    def ux(self, x):
        return (x / self.W - 0.5) * FW

    def uy(self, y):
        return (0.5 - y / self.H) * FH

    def uw(self, e):
        return e / self.W * FW

    def uh(self, e):
        return e / self.H * FH

    def theme(self, slide) -> _Theme:
        master = slide.slide_layout.slide_master
        key = id(master)
        if key not in self.themes:
            self.themes[key] = _Theme(master, self.qn)
        return self.themes[key]

    def slide_bg(self, slide, theme: _Theme) -> str:
        qn = self.qn
        for holder in (slide, slide.slide_layout, slide.slide_layout.slide_master):
            try:
                bg = holder._element.cSld.bg
            except Exception:                                   # noqa: BLE001
                bg = None
            if bg is None:
                continue
            bgpr = bg.find(qn("p:bgPr"))
            if bgpr is not None:
                for child in bgpr:
                    tag = _local(child)
                    if tag in _FILL_TAGS:
                        c = theme.fill(tag, child)
                        if c:
                            return c[0]
            ref = bg.find(qn("p:bgRef"))
            if ref is not None:
                c = theme.first_clr(ref)
                if c:
                    return c[0]
        return "#" + theme.colors.get(theme.clr_map.get("bg1", "lt1"), "FFFFFF")

    @staticmethod
    def _fill_child(parent):
        if parent is None:
            return None, None
        for c in parent:
            t = _local(c)
            if t in _FILL_TAGS:
                return t, c
        return None, None

    def shape_fill(self, spPr, style, theme: _Theme, ctx):
        tag, el = self._fill_child(spPr)
        if tag == "noFill":
            return None
        if tag == "grpFill":
            return ctx.get("group_fill")
        if tag:
            return theme.fill(tag, el)
        if spPr is not None and spPr.getparent() is not None and spPr.getparent().get("useBgFill") == "1":
            return (ctx["bg"], 1.0)
        if style is not None:
            ref = style.find(self.qn("a:fillRef"))
            if ref is not None and _int(ref.get("idx")) > 0:
                return theme.first_clr(ref)
        return None

    def shape_line(self, spPr, style, theme: _Theme):
        """{"color", "alpha", "w" (EMU), "dash", "head", "tail"} or None when the shape has no outline."""
        qn = self.qn
        ln = spPr.find(qn("a:ln")) if spPr is not None else None
        color = None
        w = None
        have_fill = False
        if ln is not None:
            w = _int(ln.get("w"), 0) or None
            tag, el = self._fill_child(ln)
            if tag == "noFill":
                return None
            if tag:
                color = theme.fill(tag, el)
                have_fill = True
        if not have_fill:
            if style is None:
                return None
            ref = style.find(qn("a:lnRef"))
            idx = _int(ref.get("idx")) if ref is not None else 0
            if idx <= 0:
                return None
            color = theme.first_clr(ref)
            if w is None:
                w = theme.line_widths[min(idx, len(theme.line_widths) - 1)]
        if color is None:
            return None
        out = {"color": color[0], "alpha": color[1], "w": w or 9525, "dash": None, "head": "none", "tail": "none"}
        if ln is not None:
            d = ln.find(qn("a:prstDash"))
            if d is not None:
                out["dash"] = d.get("val")
            for k in ("headEnd", "tailEnd"):
                e = ln.find(qn("a:" + k))
                if e is not None:
                    out[k[:4]] = e.get("type", "none")
        return out

    def stroke_px(self, w_emu: float) -> float:
        return max(0.5, w_emu * self.px_per_emu / 0.9)

    def style_kw(self, fill, line) -> str:
        parts = []
        if line:
            parts.append(f'color="{line["color"]}", stroke_width={_num(self.stroke_px(line["w"]), 1)}')
            if line["alpha"] < 0.999:
                parts.append(f"stroke_opacity={_num(line['alpha'], 2)}")
        else:
            parts.append("stroke_width=0")
        if fill:
            parts.append(f'fill_color="{fill[0]}", fill_opacity={_num(fill[1], 2)}')
        return ", ".join(parts)

    # ----------------------------------------------------------------------------------------------- geometry
    def box_of(self, xfrm, xf: _Xf) -> _Box | None:
        qn = self.qn
        if xfrm is None:
            return None
        off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
        if off is None or ext is None:
            return None
        l, t = _int(off.get("x")), _int(off.get("y"))
        w, h = _int(ext.get("cx")), _int(ext.get("cy"))
        cx, cy = xf.pt(l + w / 2, t + h / 2)
        rot = _int(xfrm.get("rot")) / 60000
        return _Box(cx, cy, w * xf.scale_x, h * xf.scale_y, xf.angle(rot),
                    xfrm.get("flipH") == "1", xfrm.get("flipV") == "1", (l, t, w, h))

    def rot_code(self, angle_deg: float) -> str:
        return f".rotate({_num(-math.radians(angle_deg), 4)})" if abs(angle_deg) > 0.05 else ""

    def emit_shape(self, shape, xf: _Xf, ctx: dict):
        qn = self.qn
        el = shape._element
        tag = _local(el)
        if tag == "grpSp":
            grp = el.find(qn("p:grpSpPr"))
            child_xf = _Xf(xf, grp.find(qn("a:xfrm")) if grp is not None else None, qn)
            sub = dict(ctx)
            ftag, fel = self._fill_child(grp)
            if ftag and ftag not in ("grpFill",):
                sub["group_fill"] = None if ftag == "noFill" else ctx["theme"].fill(ftag, fel)
            for child in self.iter_shapes(shape.shapes, shape.shapes._spTree):
                self.emit_shape(child, child_xf, sub)
            return
        if tag == "pic":
            self.emit_picture(shape, el, xf, ctx)
        elif tag == "graphicFrame":
            self.emit_frame(shape, el, xf, ctx)
        elif tag in ("sp", "cxnSp"):
            self.emit_sp(shape, el, xf, ctx)

    # ----------------------------------------------------------------------------------------------- pictures
    def emit_picture(self, shape, el, xf: _Xf, ctx: dict):
        qn = self.qn
        box = self.box_of(self.shape_xfrm(el, ctx["slide"]), xf)
        if box is None:
            return
        try:
            blob, ext = shape.image.blob, shape.image.ext
        except Exception:                                       # noqa: BLE001
            self.warnings.append(f"slide {ctx['slide_idx']}: a picture could not be read")
            return
        fill = el.find(qn("p:blipFill"))
        blip = fill.find(qn("a:blip")) if fill is not None else None
        src = fill.find(qn("a:srcRect")) if fill is not None else None
        self.emit_blip(blob, ext, blip, src, el.find(qn("p:spPr")), el.find(qn("p:style")), box, ctx)

    def emit_blip(self, blob: bytes, ext: str, blip, src_rect, spPr, style, box: _Box, ctx: dict):
        """A picture (or a shape filled with one) as an ImageMobject: cropped, recoloured for a transparent
        colour, clipped to a rounded or elliptic frame, then its outline if it has one."""
        qn = self.qn
        si = ctx["slide_idx"]
        theme: _Theme = ctx["theme"]
        crop = tuple(_int(src_rect.get(k)) / 100000 for k in ("l", "t", "r", "b")) if src_rect is not None else (0, 0, 0, 0)
        transparent, opacity = None, 1.0
        if blip is not None:
            cc = blip.find(qn("a:clrChange"))
            if cc is not None:
                c = theme.first_clr(cc.find(qn("a:clrFrom")))
                transparent = c[0] if c else None
            am = blip.find(qn("a:alphaModFix"))
            if am is not None:
                opacity = _int(am.get("amt"), 100000) / 100000
        prst, adj = "rect", {}
        if spPr is not None:
            pg = spPr.find(qn("a:prstGeom"))
            if pg is not None:
                prst = pg.get("prst") or "rect"
                for gd in pg.iter(qn("a:gd")):
                    m = re.match(r"val\s+(-?\d+)", gd.get("fmla") or "")
                    if m:
                        adj[gd.get("name")] = int(m.group(1))
        geometry = (prst, adj) if prst in ("roundRect", "ellipse") else None
        result = self.picture_bytes(blob, ext, crop, box.flip_h, box.flip_v, transparent, geometry)
        if result is None:
            self.warnings.append(f"slide {si}: a {ext} picture could not be converted and was left out")
            return
        blob, ext = result
        self.n_img += 1
        name = f"slide{si:02d}_img{self.n_img}.{ext}"
        (self.assets / name).write_bytes(blob)
        w_u, h_u = self.uw(box.w), self.uh(box.h)
        fit = f".set(width={_num(w_u)})"
        try:
            from PIL import Image
            iw, ih = Image.open(io.BytesIO(blob)).size
            if iw and ih and abs((w_u / h_u) / (iw / ih) - 1) > 0.03:
                fit += f".stretch_to_fit_height({_num(h_u)})"
        except Exception:                                       # noqa: BLE001
            pass
        if opacity < 0.999:
            fit += f".set_opacity({_num(opacity, 2)})"
        ctx["lines"].append(f'self.add(ImageMobject(ASSETS / "{name}"){fit}.move_to(({_num(self.ux(box.cx))}, '
                            f'{_num(self.uy(box.cy))}, 0)){self.rot_code(box.angle)})  # manimble')
        line = self.shape_line(spPr, style, theme) if spPr is not None else None
        if line:
            frame = prst if prst in self.DASHABLE or prst == "ellipse" else "rect"
            self.emit_geometry(frame, adj, box, None, line, spPr, ctx["lines"], si)

    @staticmethod
    def picture_bytes(blob: bytes, ext: str, crop, flip_h: bool, flip_v: bool, transparent=None, geometry=None):
        """The picture as shown: cropped (srcRect), with its transparent colour removed, clipped to its frame
        and flipped. The original bytes when nothing applies; None when the format cannot be read."""
        plain = not any(c > 0.001 for c in crop) and not flip_h and not flip_v and transparent is None and geometry is None
        if plain and ext.lower() in ("png", "jpg", "jpeg", "gif", "bmp", "tif", "tiff", "webp"):
            return blob, ext
        try:
            from PIL import Image, ImageChops, ImageDraw
            im = Image.open(io.BytesIO(blob))
            im.load()
        except Exception:                                       # noqa: BLE001
            return None
        if plain:
            return blob, ext
        try:
            w, h = im.size
            l, t, r, b = crop
            x0, y0, x1, y1 = round(l * w), round(t * h), round(w - r * w), round(h - b * h)
            if x1 - x0 >= 1 and y1 - y0 >= 1 and (x0, y0, x1, y1) != (0, 0, w, h):
                if min(x0, y0) < 0 or x1 > w or y1 > h:
                    # a negative crop pads the picture: it sits inside a larger transparent frame
                    canvas = Image.new("RGBA", (x1 - x0, y1 - y0), (0, 0, 0, 0))
                    canvas.paste(im.convert("RGBA"), (-x0, -y0))
                    im = canvas
                else:
                    im = im.crop((x0, y0, x1, y1))
            if transparent is not None:
                import numpy as np
                im = im.convert("RGBA")
                arr = np.array(im)
                rgb = tuple(int(transparent[i:i + 2], 16) for i in (1, 3, 5))
                mask = (np.abs(arr[..., :3].astype(int) - rgb) <= 12).all(axis=-1)
                arr[mask, 3] = 0
                im = Image.fromarray(arr)
            if geometry is not None:
                prst, adj = geometry
                im = im.convert("RGBA")
                mask = Image.new("L", im.size, 0)
                d = ImageDraw.Draw(mask)
                if prst == "ellipse":
                    d.ellipse((0, 0, im.width - 1, im.height - 1), fill=255)
                else:
                    r = min(im.width, im.height) * adj.get("adj", 16667) / 100000
                    d.rounded_rectangle((0, 0, im.width - 1, im.height - 1), radius=r, fill=255)
                im.putalpha(ImageChops.multiply(im.getchannel("A"), mask))
            if flip_h:
                im = im.transpose(Image.FLIP_LEFT_RIGHT)
            if flip_v:
                im = im.transpose(Image.FLIP_TOP_BOTTOM)
            out = io.BytesIO()
            if ext.lower() in ("jpg", "jpeg") and im.mode in ("RGB", "L"):
                im.save(out, "JPEG", quality=92)
                return out.getvalue(), "jpg"
            im.save(out, "PNG")
            return out.getvalue(), "png"
        except Exception:                                       # noqa: BLE001
            return blob, ext

    # ------------------------------------------------------------------------------------- tables and charts
    def emit_frame(self, shape, el, xf: _Xf, ctx: dict):
        qn = self.qn
        si = ctx["slide_idx"]
        box = self.box_of(el.find(qn("p:xfrm")), xf)
        if box is None:
            return
        if getattr(shape, "has_table", False) and shape.has_table:
            self.emit_table(shape, box, ctx)
            return
        if getattr(shape, "has_chart", False) and shape.has_chart:
            self.emit_chart(shape, box, ctx)
            return
        gd = el.find(".//" + qn("a:graphicData"))
        kind = (gd.get("uri") or "").rsplit("/", 1)[-1] if gd is not None else "object"
        self.warnings.append(f"slide {si}: a {kind} object was not imported")

    def emit_table(self, shape, box: _Box, ctx: dict):
        """manim's Table sizes each column to its text, which collapses empty columns; lay the cells out on
        PowerPoint's column widths and row heights instead (table_grid in common.py)."""
        qn = self.qn
        theme: _Theme = ctx["theme"]
        PP_ALIGN, MSO_ANCHOR = self.PP_ALIGN, self.MSO_ANCHOR
        tbl = shape.table
        rows, sizes, fonts, colors, haligns, valigns = [], [], [], [], [], []
        for r in tbl.rows:
            rows.append([c.text.strip().replace("\n", " ") or " " for c in r.cells])
            for c in r.cells:
                if not c.text.strip():
                    continue
                para = c.text_frame.paragraphs[0]
                rpr = next((rp for t, rp, _ in self._runs(para._p) if rp is not None and t.strip()), None)
                sz = rpr.get("sz") if rpr is not None else None
                sizes.append(_int(sz) / 100 if sz else 18)
                lat = rpr.find(qn("a:latin")) if rpr is not None else None
                fonts.append(theme.font(lat.get("typeface")) if lat is not None else theme.fonts["minor"])
                col = theme.first_clr(rpr.find(qn("a:solidFill"))) if rpr is not None and rpr.find(qn("a:solidFill")) is not None else None
                colors.append(col[0] if col else ctx["default_text"])
                haligns.append(para.alignment)
                valigns.append(c.vertical_anchor)

        def most(values, default):
            return max(set(values), key=values.count) if values else default
        size_pt = most(sizes, 18)
        fs = max(6.0, size_pt * self.s_ref / PT_PER_FONT_SIZE)
        # PowerPoint's default cell insets (0.1in, 0.05in); the cap height starts about a quarter em below the inset
        margin = f"({_num(self.uw(91440))}, {_num(self.uw(45720 + 0.25 * size_pt * 12700))})"
        h = {PP_ALIGN.CENTER: "", PP_ALIGN.RIGHT: "RIGHT"}.get(most(haligns, None), "LEFT")
        v = {MSO_ANCHOR.MIDDLE: "", MSO_ANCHOR.BOTTOM: "DOWN"}.get(most(valigns, None), "UP")
        align = " + ".join(d for d in (v, h) if d) or "ORIGIN"
        kx = box.w / max(1, sum(c.width for c in tbl.columns))        # the frame may be scaled by a group
        col_w = [_num(self.uw(c.width * kx)) for c in tbl.columns]
        row_h = [_num(self.uh(r.height)) for r in tbl.rows]
        font = most(fonts, theme.fonts["minor"])
        self.fonts_used.add(font)
        self.uses_tables = True
        color = most(colors, ctx["default_text"])
        left, top = box.cx - box.w / 2, box.cy - box.h / 2
        ctx["lines"].append(
            f'self.add(table_grid(Table({rows!r}, include_outer_lines=True, line_config={{"color": "{ctx["default_text"]}"}}, '
            f'element_to_mobject_config={{"color": "{color}", "font": "{font}", "font_size": {_num(fs, 1)}}}), '
            f'[{", ".join(col_w)}], [{", ".join(row_h)}], align={align}, margin={margin})'
            f'.move_to(({_num(self.ux(left))}, {_num(self.uy(top))}, 0), UL))  # manimble')

    CHART_NS = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"

    def emit_chart(self, shape, box: _Box, ctx: dict):
        """Charts are drawn from their data with plain manim shapes (bars, lines, sectors, labels): the layout
        follows PowerPoint's defaults rather than the chart's own formatting."""
        C = self.CHART_NS
        qn = self.qn
        si = ctx["slide_idx"]
        theme: _Theme = ctx["theme"]
        out = ctx["lines"]
        try:
            cs = shape.chart._chartSpace
            chart_el = cs.find(C + "chart")
            plot_area = chart_el.find(C + "plotArea")
        except Exception as exc:                                # noqa: BLE001
            self.warnings.append(f"slide {si}: a chart could not be read ({exc})")
            return

        def val(e, tag, default=None):
            if e is None:
                return default
            c = e.find(C + tag)
            return c.get("val", default) if c is not None else default

        def texts(e):
            if e is None:
                return ""
            s = "".join(t.text or "" for t in e.iter(qn("a:t")))
            return s or " ".join(v.text or "" for v in e.iter(C + "v"))

        def points(e):
            if e is None:
                return []
            pts = {}
            for pt in e.iter(C + "pt"):
                v = pt.find(C + "v")
                pts[_int(pt.get("idx"))] = v.text if v is not None else None
            return [pts.get(i) for i in range(max(pts) + 1)] if pts else []

        def numbers(e):
            vals = []
            for v in points(e):
                try:
                    vals.append(float(v))
                except (TypeError, ValueError):
                    vals.append(None)
            return vals

        palette = ["#" + theme.colors.get(f"accent{i}", "4472C4") for i in range(1, 7)]
        label_color = "#595959"
        font = theme.fonts["minor"]
        self.fonts_used.add(font)
        fs_pt = 10.0
        fs = fs_pt * self.s_ref / PT_PER_FONT_SIZE
        pil = _font_or_fallback(font, False, False)

        def text_w(s, size):                                    # width of a label in units
            ems = pil.getlength(s) / 1000 if pil is not None else len(s) * 0.5
            return ems * size / 72

        cats: list = []
        plots, n_series = [], 0
        for pe in plot_area:
            tag = _local(pe)
            if not tag.endswith("Chart"):
                continue
            kind = tag[:-5]
            grouping = val(pe, "grouping", "clustered" if kind == "bar" else "standard")
            sers = []
            for se in pe.findall(C + "ser"):
                name = texts(se.find(C + "tx")) or f"Series {n_series + 1}"
                vals = numbers(se.find(C + "yVal") if kind in ("scatter", "bubble") else se.find(C + "val"))
                xs = numbers(se.find(C + "xVal")) if kind in ("scatter", "bubble") else []
                c = se.find(C + "cat")
                if c is not None and not cats:
                    cats = [str(v) if v is not None else "" for v in points(c)]
                spPr = se.find(C + "spPr")
                color = None
                if spPr is not None:
                    ftag, fel = self._fill_child(spPr)
                    if ftag:
                        cc = theme.fill(ftag, fel)
                        color = cc[0] if cc else None
                    ln = spPr.find(qn("a:ln"))
                    if color is None and ln is not None:
                        ltag, lel = self._fill_child(ln)
                        if ltag:
                            cc = theme.fill(ltag, lel)
                            color = cc[0] if cc else None
                pt_colors = {}
                for dpt in se.findall(C + "dPt"):
                    dsp = dpt.find(C + "spPr")
                    ftag, fel = self._fill_child(dsp)
                    cc = theme.fill(ftag, fel) if ftag else None
                    if cc:
                        pt_colors[_int(val(dpt, "idx"))] = cc[0]
                marker = val(se.find(C + "marker"), "symbol", None)
                sers.append({"name": name, "vals": vals, "xs": xs, "color": color or palette[n_series % 6],
                             "smooth": val(se, "smooth", "0") == "1", "pt_colors": pt_colors, "marker": marker})
                n_series += 1
            plots.append({"kind": kind, "grouping": grouping, "dir": val(pe, "barDir", "col"),
                          "gap": _int(val(pe, "gapWidth", 150)), "sers": sers,
                          "hole": _int(val(pe, "holeSize", 50)), "style": val(pe, "scatterStyle", "lineMarker"),
                          "first_angle": _int(val(pe, "firstSliceAng", 0))})
        if not plots or not any(p["sers"] for p in plots):
            self.warnings.append(f"slide {si}: a chart had no data and was left out")
            return

        # ------------------------------------------------------------------ frame: title, legend, plot area
        w_u, h_u = self.uw(box.w), self.uh(box.h)
        left, top = self.ux(box.cx) - w_u / 2, self.uy(box.cy) + h_u / 2
        title_el = chart_el.find(C + "title")
        title = texts(title_el) if title_el is not None and val(chart_el, "autoTitleDeleted", "0") != "1" else ""
        if title_el is not None and not title and n_series == 1 and val(chart_el, "autoTitleDeleted", "0") != "1":
            title = plots[0]["sers"][0]["name"]
        legend_el = chart_el.find(C + "legend")
        legend_pos = val(legend_el, "legendPos", "r") if legend_el is not None else None
        pies = all(p["kind"] in ("pie", "doughnut", "pie3D", "ofPie") for p in plots)
        cat_ax = plot_area.find(C + "catAx") if plot_area.find(C + "catAx") is not None else plot_area.find(C + "dateAx")
        val_axes = plot_area.findall(C + "valAx")
        val_ax = val_axes[0] if val_axes else None
        if plots[0]["kind"] == "scatter" and len(val_axes) > 1:
            cat_ax, val_ax = val_axes[0], val_axes[1]
        cat_hidden = cat_ax is None or val(cat_ax, "delete", "0") == "1"
        val_hidden = val_ax is None or val(val_ax, "delete", "0") == "1"
        gridlines = val_ax is not None and val_ax.find(C + "majorGridlines") is not None
        cat_title = texts(cat_ax.find(C + "title")) if cat_ax is not None else ""
        val_title = texts(val_ax.find(C + "title")) if val_ax is not None else ""

        x0, x1 = left + 0.03 * w_u, left + 0.97 * w_u
        y_top, y_bot = top - 0.04 * h_u, top - 0.96 * h_u
        title_fs = fs * 1.4
        if title:
            out.append(f'self.add(Text({title!r}, font_size={_num(title_fs, 1)}, color="{label_color}", font="{font}").move_to(({_num((x0 + x1) / 2)}, {_num(y_top - 0.08 * h_u)}, 0)))  # manimble')
            y_top -= 0.16 * h_u
        entries = [(s["color"], s["name"]) for p in plots for s in p["sers"]]
        if pies:
            entries = [(plots[0]["sers"][0]["pt_colors"].get(i, palette[i % 6]), c) for i, c in enumerate(cats)]
        legend_h = 0.1 * h_u
        if legend_el is not None and legend_pos in ("t", "b", "tr") and entries:
            ly = y_top - legend_h / 2 if legend_pos != "b" else y_bot + legend_h / 2
            sq = fs / 72 * 0.8
            widths = [sq * 1.6 + text_w(n, fs) + 0.25 for _, n in entries]
            x = (x0 + x1) / 2 - sum(widths) / 2
            for (col, name), wi in zip(entries, widths):
                out.append(f'self.add(Square(side_length={_num(sq)}, fill_color="{col}", fill_opacity=1, stroke_width=0).move_to(({_num(x + sq / 2)}, {_num(ly)}, 0)))  # manimble')
                out.append(f'self.add(Text({name!r}, font_size={_num(fs, 1)}, color="{label_color}", font="{font}").move_to(({_num(x + sq * 1.6 + text_w(name, fs) / 2)}, {_num(ly)}, 0)))  # manimble')
                x += wi
            if legend_pos == "b":
                y_bot += legend_h
            else:
                y_top -= legend_h
        elif legend_el is not None and legend_pos in ("r", "l") and entries:
            lw = max(text_w(n, fs) for _, n in entries) + fs / 72 * 1.6 + 0.1
            lx = x1 - lw if legend_pos == "r" else x0
            sq = fs / 72 * 0.8
            step = fs / 72 * 1.6
            ly = (y_top + y_bot) / 2 + step * (len(entries) - 1) / 2
            for col, name in entries:
                out.append(f'self.add(Square(side_length={_num(sq)}, fill_color="{col}", fill_opacity=1, stroke_width=0).move_to(({_num(lx + sq / 2)}, {_num(ly)}, 0)))  # manimble')
                out.append(f'self.add(Text({name!r}, font_size={_num(fs, 1)}, color="{label_color}", font="{font}").move_to(({_num(lx + sq * 1.6 + text_w(name, fs) / 2)}, {_num(ly)}, 0)))  # manimble')
                ly -= step
            if legend_pos == "r":
                x1 -= lw + 0.05 * w_u
            else:
                x0 += lw + 0.05 * w_u

        # ----------------------------------------------------------------------------------------- pies
        if pies:
            p = plots[0]
            s = p["sers"][0]
            vals = [v or 0.0 for v in s["vals"]]
            total = sum(vals) or 1.0
            r = min(x1 - x0, y_top - y_bot) * 0.45
            cx, cy = (x0 + x1) / 2, (y_top + y_bot) / 2
            start = math.pi / 2 - math.radians(p["first_angle"])
            hole = r * p["hole"] / 100 if p["kind"] == "doughnut" else 0.0
            for i, v in enumerate(vals):
                ang = 2 * math.pi * v / total
                col = s["pt_colors"].get(i, palette[i % 6])
                if hole:
                    out.append(f'self.add(AnnularSector(inner_radius={_num(hole)}, outer_radius={_num(r)}, start_angle={_num(start - ang, 4)}, angle={_num(ang, 4)}, arc_center=({_num(cx)}, {_num(cy)}, 0), fill_color="{col}", fill_opacity=1, stroke_width=0))  # manimble')
                else:
                    out.append(f'self.add(Sector(radius={_num(r)}, start_angle={_num(start - ang, 4)}, angle={_num(ang, 4)}, arc_center=({_num(cx)}, {_num(cy)}, 0), fill_color="{col}", fill_opacity=1, stroke_width=0))  # manimble')
                start -= ang
            return

        # ------------------------------------------------------------------------------- axes and scales
        if val_title:
            out.append(f'self.add(Text({val_title!r}, font_size={_num(fs, 1)}, color="{label_color}", font="{font}").rotate(PI / 2).move_to(({_num(x0 + fs / 72 * 0.6)}, {_num((y_top + y_bot) / 2)}, 0)))  # manimble')
            x0 += fs / 72 * 1.6
        if cat_title:
            out.append(f'self.add(Text({cat_title!r}, font_size={_num(fs, 1)}, color="{label_color}", font="{font}").move_to(({_num((x0 + x1) / 2)}, {_num(y_bot + fs / 72 * 0.6)}, 0)))  # manimble')
            y_bot += fs / 72 * 1.6
        if not cat_hidden:
            y_bot += fs / 72 * 1.5
        scatter = plots[0]["kind"] in ("scatter", "bubble")
        n_cat = max([len(s["vals"]) for p in plots for s in p["sers"]] + [len(cats), 1])
        stacked = {p["grouping"] for p in plots if p["kind"] in ("bar", "area")} & {"stacked", "percentStacked"}
        lo, hi = 0.0, 0.0
        for p in plots:
            if p["kind"] in ("bar", "area") and p["grouping"] in ("stacked", "percentStacked"):
                for i in range(n_cat):
                    pos = sum(s["vals"][i] for s in p["sers"] if i < len(s["vals"]) and s["vals"][i] and s["vals"][i] > 0)
                    neg = sum(s["vals"][i] for s in p["sers"] if i < len(s["vals"]) and s["vals"][i] and s["vals"][i] < 0)
                    hi, lo = max(hi, pos), min(lo, neg)
                if p["grouping"] == "percentStacked":
                    hi, lo = 1.0, 0.0
            else:
                for s in p["sers"]:
                    for v in s["vals"]:
                        if v is not None:
                            hi, lo = max(hi, v), min(lo, v)
        scaling = val_ax.find(C + "scaling") if val_ax is not None else None
        v_min, v_max = val(scaling, "min"), val(scaling, "max")
        major = val(val_ax, "majorUnit")

        def nice(x):
            if x <= 0:
                return 1.0
            e = 10 ** math.floor(math.log10(x))
            f = x / e
            return e * (1 if f < 1.5 else 2 if f < 3.5 else 5 if f < 7.5 else 10)

        step = float(major) if major else nice((hi - lo) / 4 or 1.0)
        bottom = float(v_min) if v_min else math.floor(lo / step) * step
        top_v = float(v_max) if v_max else math.ceil(hi / step) * step
        if top_v <= bottom:
            top_v = bottom + step
        if not val_hidden:
            n_dec = max(0, -math.floor(math.log10(step))) if step < 1 else 0

            def fmt(v):
                return f"{v:.{n_dec}f}" if n_dec else str(int(round(v)))
            x0 += max(text_w(fmt(top_v), fs), text_w(fmt(bottom), fs)) + fs / 72 * 0.5

        def ys(v):
            return y_bot + (v - bottom) / (top_v - bottom) * (y_top - y_bot)

        x_lo, x_hi = 0.0, float(n_cat)
        if scatter:
            allx = [x for p in plots for s in p["sers"] for x in s["xs"] if x is not None]
            if allx:
                xstep = nice((max(allx) - min(allx)) / 5 or 1.0)
                x_lo, x_hi = math.floor(min(allx) / xstep) * xstep, math.ceil(max(allx) / xstep) * xstep
                if x_hi <= x_lo:
                    x_hi = x_lo + xstep

        def xs(x):
            return x0 + (x - x_lo) / (x_hi - x_lo) * (x1 - x0)

        y_zero = ys(min(max(0.0, bottom), top_v))
        if gridlines:
            v = bottom
            while v <= top_v + 1e-9:
                out.append(f'self.add(Line(({_num(x0)}, {_num(ys(v))}, 0), ({_num(x1)}, {_num(ys(v))}, 0), color="#D9D9D9", stroke_width=1))  # manimble')
                v += step
        if not val_hidden:
            v = bottom
            while v <= top_v + 1e-9:
                out.append(f'self.add(Text("{fmt(v)}", font_size={_num(fs, 1)}, color="{label_color}", font="{font}").move_to(({_num(x0 - 0.3)}, {_num(ys(v))}, 0)).align_to(({_num(x0 - fs / 72 * 0.4)}, 0, 0), RIGHT))  # manimble')
                v += step
        out.append(f'self.add(Line(({_num(x0)}, {_num(y_zero)}, 0), ({_num(x1)}, {_num(y_zero)}, 0), color="#BFBFBF", stroke_width=1.5))  # manimble')
        if scatter and not val_hidden:
            out.append(f'self.add(Line(({_num(x0)}, {_num(y_bot)}, 0), ({_num(x0)}, {_num(y_top)}, 0), color="#BFBFBF", stroke_width=1.5))  # manimble')
        cat_w = (x1 - x0) / n_cat

        # --------------------------------------------------------------------------------------- series
        for p in plots:
            sers = p["sers"]
            if p["kind"] == "bar":
                if p["dir"] != "col":
                    self.warnings.append(f"slide {si}: a horizontal bar chart was drawn with vertical bars")
                if p["grouping"] in ("stacked", "percentStacked"):
                    bw = cat_w / (1 + p["gap"] / 100)
                    for i in range(n_cat):
                        total = sum(abs(s["vals"][i]) for s in sers if i < len(s["vals"]) and s["vals"][i]) or 1.0
                        pos = neg = 0.0
                        for s in sers:
                            v = s["vals"][i] if i < len(s["vals"]) else None
                            if not v:
                                continue
                            if p["grouping"] == "percentStacked":
                                v = v / total
                            base = pos if v > 0 else neg
                            ya, yb = ys(base), ys(base + v)
                            if v > 0:
                                pos += v
                            else:
                                neg += v
                            cx = x0 + (i + 0.5) * cat_w
                            out.append(f'self.add(Rectangle(width={_num(bw)}, height={_num(abs(yb - ya))}, fill_color="{s["pt_colors"].get(i, s["color"])}", fill_opacity=1, stroke_width=0).move_to(({_num(cx)}, {_num((ya + yb) / 2)}, 0)))  # manimble')
                else:
                    nb = max(1, len(sers))
                    bw = cat_w / (nb + p["gap"] / 100)
                    for i in range(n_cat):
                        start = x0 + i * cat_w + (cat_w - nb * bw) / 2
                        for j, s in enumerate(sers):
                            v = s["vals"][i] if i < len(s["vals"]) else None
                            if v is None:
                                continue
                            ya, yb = ys(max(bottom, min(0.0, top_v))), ys(v)
                            cx = start + (j + 0.5) * bw
                            out.append(f'self.add(Rectangle(width={_num(bw)}, height={_num(abs(yb - ya))}, fill_color="{s["pt_colors"].get(i, s["color"])}", fill_opacity=1, stroke_width=0).move_to(({_num(cx)}, {_num((ya + yb) / 2)}, 0)))  # manimble')
            elif p["kind"] in ("line", "scatter", "area", "radar", "bubble"):
                for s in sers:
                    if scatter:
                        pts = [(xs(x), ys(v)) for x, v in zip(s["xs"], s["vals"]) if x is not None and v is not None]
                    else:
                        pts = [(x0 + (i + 0.5) * cat_w, ys(v)) for i, v in enumerate(s["vals"]) if v is not None]
                    if len(pts) < 2:
                        continue
                    lit = ", ".join(f"({_num(x)}, {_num(y)}, 0)" for x, y in pts)
                    sw = _num(self.stroke_px(28575), 1)
                    if p["kind"] == "area":
                        poly = f"({_num(pts[0][0])}, {_num(y_zero)}, 0), " + lit + f", ({_num(pts[-1][0])}, {_num(y_zero)}, 0)"
                        out.append(f'self.add(Polygon({poly}, fill_color="{s["color"]}", fill_opacity=0.7, stroke_width=0))  # manimble')
                        continue
                    if scatter and p["style"] in ("marker", "none"):
                        pass
                    else:
                        method = "set_points_smoothly" if s["smooth"] or p["style"] in ("smooth", "smoothMarker") else "set_points_as_corners"
                        out.append(f'self.add(VMobject(stroke_color="{s["color"]}", stroke_width={sw}).{method}([{lit}]))  # manimble')
                    if (s["marker"] and s["marker"] != "none") or (scatter and p["style"] in ("marker", "lineMarker", "smoothMarker") and s["marker"] != "none"):
                        dots = ", ".join(f'Dot(({_num(x)}, {_num(y)}, 0), radius=0.06, color="{s["color"]}")' for x, y in pts)
                        out.append(f"self.add({dots})  # manimble")
        # ------------------------------------------------------------------------------- category labels
        if not cat_hidden:
            if scatter:
                x = x_lo
                while x <= x_hi + 1e-9:
                    out.append(f'self.add(Text("{x:g}", font_size={_num(fs, 1)}, color="{label_color}", font="{font}").move_to(({_num(xs(x))}, {_num(y_bot - fs / 72 * 0.8)}, 0)))  # manimble')
                    x += xstep
            else:
                labels = cats[:n_cat] if cats else [str(i + 1) for i in range(n_cat)]
                every = max(1, math.ceil(sum(text_w(l, fs) + 0.1 for l in labels) / max(x1 - x0, 0.1)))
                for i, lab in enumerate(labels):
                    if i % every or not lab:
                        continue
                    out.append(f'self.add(Text({lab!r}, font_size={_num(fs, 1)}, color="{label_color}", font="{font}").move_to(({_num(x0 + (i + 0.5) * cat_w)}, {_num(y_bot - fs / 72 * 0.8)}, 0)))  # manimble')
        self.warnings.append(f"slide {si}: a chart was redrawn from its data with manim shapes")

    # --------------------------------------------------------------------------------- shapes & connectors
    LINE_PRSTS = ("line", "straightConnector1", "bentConnector2", "bentConnector3", "bentConnector4",
                  "bentConnector5", "curvedConnector2", "curvedConnector3", "curvedConnector4", "curvedConnector5")

    def emit_sp(self, shape, el, xf: _Xf, ctx: dict):
        qn = self.qn
        si = ctx["slide_idx"]
        theme: _Theme = ctx["theme"]
        spPr = el.find(qn("p:spPr"))
        style = el.find(qn("p:style"))
        box = self.box_of(self.shape_xfrm(el, ctx["slide"]), xf)
        if box is None:
            return
        prst_el = spPr.find(qn("a:prstGeom"))
        prst = prst_el.get("prst") if prst_el is not None else ("custom" if spPr.find(qn("a:custGeom")) is not None else "rect")
        adj = {}
        if prst_el is not None:
            for gd in prst_el.iter(qn("a:gd")):
                m = re.match(r"val\s+(-?\d+)", gd.get("fmla") or "")
                if m:
                    adj[gd.get("name")] = int(m.group(1))
        fill = self.shape_fill(spPr, style, theme, ctx)
        line = self.shape_line(spPr, style, theme)
        lines = ctx["lines"]
        txBody = el.find(qn("p:txBody"))
        has_text = txBody is not None and "".join(t.text or "" for t in txBody.iter(qn("a:t"), _MATH_NS + "t")).strip() != ""

        if prst in self.LINE_PRSTS or _local(el) == "cxnSp":
            if line:
                l, t, w, h = box.raw
                x1, x2 = (l + w, l) if box.flip_h else (l, l + w)
                y1, y2 = (t + h, t) if box.flip_v else (t, t + h)
                p1, p2 = xf.pt(x1, y1), xf.pt(x2, y2)
                self.emit_line(p1, p2, line, lines)
                if prst.startswith("curved") or prst.startswith("bent"):
                    self.warnings.append(f"slide {si}: a {prst} connector was drawn straight")
            return

        ftag, fel = self._fill_child(spPr)
        if ftag == "blipFill":                                  # a shape filled with a picture (e.g. an equation's fallback)
            blip = fel.find(qn("a:blip"))
            rid = blip.get(qn("r:embed")) if blip is not None else None
            try:
                part = ctx["slide"].part.related_part(rid)
                self.emit_blip(part.blob, part.ext, blip, fel.find(qn("a:srcRect")), spPr, style, box, ctx)
            except Exception:                                   # noqa: BLE001
                self.warnings.append(f"slide {si}: a picture fill could not be read")
        elif fill or line:
            self.emit_geometry(prst, adj, box, fill, line, spPr, lines, si)
        if has_text:
            self.emit_text(el, txBody, box, style, ctx)

    def emit_line(self, p1, p2, line: dict, lines: list):
        a = f"({_num(self.ux(p1[0]))}, {_num(self.uy(p1[1]))}, 0)"
        b = f"({_num(self.ux(p2[0]))}, {_num(self.uy(p2[1]))}, 0)"
        sw = self.stroke_px(line["w"])
        style = f'color="{line["color"]}", stroke_width={_num(sw, 1)}'
        head, tail = line["head"] != "none", line["tail"] != "none"
        if head or tail:
            # PowerPoint's medium arrowhead is about three line widths long and wide
            tip = _num(max(0.12, self.uw(line["w"]) * 3.2))
            cls = "DoubleArrow" if head and tail else "Arrow"
            if head and not tail:
                a, b = b, a
            lines.append(f"self.add({cls}({a}, {b}, buff=0, {style}, tip_length={tip}, "
                         f"max_tip_length_to_length_ratio=1, max_stroke_width_to_length_ratio=100))  # manimble")
        else:
            code = f"Line({a}, {b}, {style})"
            if line.get("dash") not in (None, "solid"):
                length = math.hypot(self.uw(p2[0] - p1[0]), self.uh(p2[1] - p1[1]))
                code = self.dashed(code, length, line)
            lines.append(f"self.add({code})  # manimble")

    def emit_geometry(self, prst, adj, box: _Box, fill, line, spPr, lines, si):
        qn = self.qn
        w_u, h_u = self.uw(box.w), self.uh(box.h)
        cx, cy = self.ux(box.cx), self.uy(box.cy)
        c = f"({_num(cx)}, {_num(cy)}, 0)"
        dashed = bool(line and line.get("dash") not in (None, "solid")) and prst in self.DASHABLE
        if dashed and fill:
            # manim dashes a stroke by cutting the shape into pieces, which would cut up the fill too: the fill
            # is its own shape underneath the dashed outline
            self.emit_geometry(prst, adj, box, fill, None, spPr, lines, si)
            fill = None
        kw = self.style_kw(fill, line)
        rc = self.rot_code(box.angle)
        ss = min(w_u, h_u)

        def poly(points, closed=True, filled=True):
            # points in box coordinates (fractions of w, h from the top-left corner) -> absolute units
            pts = []
            for px, py in points:
                x, y = (px - 0.5) * w_u, (0.5 - py) * h_u
                if box.flip_h:
                    x = -x
                if box.flip_v:
                    y = -y
                pts.append(f"({_num(cx + x)}, {_num(cy + y)}, 0)")
            rot = f".rotate({_num(-math.radians(box.angle), 4)}, about_point={c})" if abs(box.angle) > 0.05 else ""
            if closed:
                return f"Polygon({', '.join(pts)}, {kw if filled else self.style_kw(None, line)}){rot}"
            stroke = f'stroke_color="{line["color"]}", stroke_width={_num(self.stroke_px(line["w"]), 1)}' if line else "stroke_width=0"
            return f"VMobject({stroke}).set_points_as_corners([{', '.join(pts)}]){rot}"

        if prst == "ellipse":
            code = f"Ellipse(width={_num(w_u)}, height={_num(h_u)}, {kw}).move_to({c}){rc}"
        elif prst in ("roundRect", "round1Rect", "round2SameRect", "round2DiagRect", "snipRoundRect"):
            r = ss * adj.get("adj", adj.get("adj1", 16667)) / 100000
            code = f"RoundedRectangle(width={_num(w_u)}, height={_num(h_u)}, corner_radius={_num(max(0.001, min(r, ss / 2)))}, {kw}).move_to({c}){rc}"
        elif prst == "triangle":
            code = f"Triangle({kw}).stretch_to_fit_width({_num(w_u)}).stretch_to_fit_height({_num(h_u)}).move_to({c}){rc}"
        elif prst == "diamond":
            code = poly([(0.5, 0), (1, 0.5), (0.5, 1), (0, 0.5)])
        elif prst in ("rightArrow", "leftArrow", "upArrow", "downArrow"):
            horizontal = prst in ("rightArrow", "leftArrow")
            L, T = (w_u, h_u) if horizontal else (h_u, w_u)      # along the arrow, across it
            a1 = min(max(adj.get("adj1", 50000), 0), 100000) / 100000
            a2 = min(max(adj.get("adj2", 50000), 0), 100000 * L / ss) / 100000
            head = ss * a2 / L if L else 0.5                     # fraction of the length
            half = a1 / 2
            pts = [(0, 0.5 - half), (1 - head, 0.5 - half), (1 - head, 0), (1, 0.5), (1 - head, 1), (1 - head, 0.5 + half), (0, 0.5 + half)]
            if prst == "leftArrow":
                pts = [(1 - x, y) for x, y in pts]
            elif prst == "downArrow":
                pts = [(y, x) for x, y in pts]
            elif prst == "upArrow":
                pts = [(y, 1 - x) for x, y in pts]
            code = poly(pts)
        elif prst == "pie":
            a1, a2 = adj.get("adj1", 0) / 60000, adj.get("adj2", 16200000) / 60000
            sweep = (a2 - a1) % 360 or 360
            code = (f"Sector(radius={_num(w_u / 2)}, start_angle={_num(-math.radians(a2), 4)}, angle={_num(math.radians(sweep), 4)}, "
                    f"arc_center={c}, {kw})")
            if abs(w_u - h_u) > 1e-3:
                code += f".stretch({_num(h_u / w_u, 4)}, 1, about_point={c})"
            code += rc
        elif prst in ("leftBrace", "rightBrace"):
            # PowerPoint's brace: two hooks and a tip made of quarter ellipses, drawn as a stroked path
            if not line:
                line = {"color": fill[0] if fill else "#000000", "w": 9525, "alpha": 1.0, "dash": None}
            y1 = min(0.5, ss * adj.get("adj1", 8333) / 100000 / h_u) if h_u else 0.1
            y3 = min(1 - y1, max(y1, adj.get("adj2", 50000) / 100000))
            pts = []

            def arc(ccx, ccy, a0, a1):
                for k in range(9):
                    a = math.radians(a0 + (a1 - a0) * k / 8)
                    pts.append((ccx + 0.5 * math.cos(a), ccy + y1 * math.sin(a)))
            arc(0, y1, -90, 0)
            arc(1, y3 - y1, 180, 90)
            arc(1, y3 + y1, -90, -180)
            arc(0, 1 - y1, 0, 90)
            if prst == "leftBrace":
                pts = [(1 - x, yv) for x, yv in pts]
            code = poly(pts, closed=False)
        elif prst == "noSmoking":
            thick = ss * adj.get("adj", 18750) / 100000
            r_out, r_in = w_u / 2, max(0.01, w_u / 2 - thick)
            ring_kw = kw if fill else kw.replace("stroke_width=0", "stroke_width=0")
            code = f"Annulus(inner_radius={_num(r_in)}, outer_radius={_num(r_out)}, {ring_kw}).move_to({c})"
            if abs(w_u - h_u) > 1e-3:
                code += f".stretch({_num(h_u / w_u, 4)}, 1, about_point={c})"
            lines.append(f"self.add({code}{rc})  # manimble")
            bar_kw = f'fill_color="{fill[0]}", fill_opacity={_num(fill[1], 2)}, stroke_width=0' if fill else kw
            code = f"Rectangle(width={_num(2 * r_in * 0.98)}, height={_num(thick)}, {bar_kw}).rotate(-PI / 4).move_to({c}){rc}"
        elif prst == "custom":
            paths = self.custom_paths(spPr, box)
            if not paths:
                code = f"Rectangle(width={_num(w_u)}, height={_num(h_u)}, {kw}).move_to({c}){rc}"
            else:
                for closed, pts, filled in paths[:-1]:
                    lines.append(f"self.add({poly(pts, closed, filled)})  # manimble")
                closed, pts, filled = paths[-1]
                code = poly(pts, closed, filled)
        else:
            if prst not in ("rect", "flowChartProcess", "plaque", "snip1Rect", "snip2SameRect", "snip2DiagRect"):
                self.warnings.append(f"slide {si}: a {prst} shape was imported as a rectangle")
            code = f"Rectangle(width={_num(w_u)}, height={_num(h_u)}, {kw}).move_to({c}){rc}"
        if dashed:
            if prst == "ellipse":
                a, b = w_u / 2, h_u / 2
                perimeter = math.pi * (3 * (a + b) - math.sqrt((3 * a + b) * (a + 3 * b)))
            else:
                perimeter = 2 * (w_u + h_u)
            code = self.dashed(code, perimeter, line)
        lines.append(f"self.add({code})  # manimble")

    DASHABLE = ("rect", "roundRect", "round1Rect", "round2SameRect", "round2DiagRect", "snipRoundRect", "ellipse",
                "triangle", "diamond", "rightArrow", "leftArrow", "upArrow", "downArrow", "flowChartProcess")
    # PowerPoint's preset dashes as (dash, gap) in line widths
    DASHES = {"dash": (4, 3), "sysDash": (3, 1), "dot": (1, 3), "sysDot": (1, 1), "lgDash": (8, 3), "dashDot": (4, 3),
              "lgDashDot": (8, 3), "lgDashDotDot": (8, 3), "sysDashDot": (3, 1), "sysDashDotDot": (3, 1)}

    def dashed(self, code: str, length_u: float, line: dict) -> str:
        d, g = self.DASHES.get(line.get("dash") or "dash", (4, 3))
        lw = max(self.uw(line["w"]), 0.01)
        n = max(2, round(length_u / ((d + g) * lw)))
        return f"DashedVMobject({code}, num_dashes={n}, dashed_ratio={_num(d / (d + g), 2)})"

    def custom_paths(self, spPr, box: _Box):
        """[(closed, [(fx, fy)], filled)] for a custGeom: each sub-path as points in fractions of the box,
        curves sampled."""
        qn = self.qn
        out = []
        for path in spPr.iter(qn("a:path")):
            pw, ph = float(path.get("w") or 0) or 1.0, float(path.get("h") or 0) or 1.0
            filled = path.get("fill", "norm") != "none"
            cur = None
            pts: list = []

            def p(e):
                return _int(e.get("x")) / pw, _int(e.get("y")) / ph

            def flush(closed):
                if len(pts) >= 2:
                    out.append((closed, list(pts), filled))
                pts.clear()

            for cmd in path:
                t = _local(cmd)
                if t == "moveTo":
                    flush(False)
                    cur = p(cmd[0])
                    pts.append(cur)
                elif t == "lnTo":
                    cur = p(cmd[0])
                    pts.append(cur)
                elif t in ("cubicBezTo", "quadBezTo") and cur is not None:
                    cps = [p(e) for e in cmd]
                    if t == "quadBezTo" and len(cps) == 2:
                        q, e = cps
                        cps = [(cur[0] + 2 / 3 * (q[0] - cur[0]), cur[1] + 2 / 3 * (q[1] - cur[1])),
                               (e[0] + 2 / 3 * (q[0] - e[0]), e[1] + 2 / 3 * (q[1] - e[1])), e]
                    if len(cps) == 3:
                        (x1, y1), (x2, y2), (x3, y3) = cps
                        x0, y0 = cur
                        for k in range(1, 9):
                            s = k / 8
                            u = 1 - s
                            pts.append((u ** 3 * x0 + 3 * u * u * s * x1 + 3 * u * s * s * x2 + s ** 3 * x3,
                                        u ** 3 * y0 + 3 * u * u * s * y1 + 3 * u * s * s * y2 + s ** 3 * y3))
                        cur = cps[-1]
                elif t == "arcTo" and cur is not None:
                    wr, hr = _int(cmd.get("wR")) / pw, _int(cmd.get("hR")) / ph
                    st, sw = _int(cmd.get("stAng")) / 60000, _int(cmd.get("swAng")) / 60000
                    ccx = cur[0] - wr * math.cos(math.radians(st))
                    ccy = cur[1] - hr * math.sin(math.radians(st))
                    for k in range(1, 13):
                        a = math.radians(st + sw * k / 12)
                        cur = (ccx + wr * math.cos(a), ccy + hr * math.sin(a))
                        pts.append(cur)
                elif t == "close":
                    flush(True)
            flush(False)
        return out

    # ----------------------------------------------------------------------------------------------- text
    def _runs(self, p_el):
        """[(text, rPr element or None, math element or None)] for the runs, fields, line breaks and Office Math
        of a paragraph. Math comes with a plain-text rendering as its text."""
        qn = self.qn
        out = []
        for child in p_el:
            tag = _local(child)
            if tag in ("r", "fld"):
                t = child.find(qn("a:t"))
                out.append(((t.text or "") if t is not None else "", child.find(qn("a:rPr")), None))
            elif tag == "br":
                out.append(("\n", child.find(qn("a:rPr")), None))
            elif tag in ("m", "oMath", "oMathPara"):
                om = child if tag != "m" else (child.find(_MATH_NS + "oMathPara") if child.find(_MATH_NS + "oMathPara") is not None
                                              else child.find(_MATH_NS + "oMath"))
                if om is None:
                    om = child
                out.append((omml_text(om), child.find(".//" + qn("a:rPr")), om))
            elif tag == "AlternateContent":
                branch = self.alternate_branch(child)
                if branch is not None:
                    out += self._runs(branch)
        return out

    @staticmethod
    def alternate_branch(ac):
        """The branch of an mc:AlternateContent to import: the Choice when it is Office Math (which is
        converted), otherwise the Fallback (usually a picture of the content)."""
        choice = next((c for c in ac if _local(c) == "Choice"), None)
        fallback = next((c for c in ac if _local(c) == "Fallback"), None)
        if choice is not None and choice.get("Requires") == "a14" and choice.find(".//" + _MATH_NS + "oMath") is not None:
            return choice
        return fallback if fallback is not None else choice

    SHAPE_TAGS = ("sp", "grpSp", "graphicFrame", "cxnSp", "pic")

    def iter_shapes(self, shapes, container):
        """python-pptx shapes for the children of a spTree / grpSp, in document order, looking inside
        mc:AlternateContent (which python-pptx skips)."""
        for child in container:
            tag = _local(child)
            if tag in self.SHAPE_TAGS:
                yield shapes._shape_factory(child)
            elif tag == "AlternateContent":
                branch = self.alternate_branch(child)
                if branch is not None:
                    yield from self.iter_shapes(shapes, branch)

    def placeholder_sources(self, el, slide) -> list:
        """The layout and master placeholders that a slide placeholder inherits from, nearest first."""
        qn = self.qn
        ph = el.find("./*/" + qn("p:nvPr") + "/" + qn("p:ph"))
        if ph is None:
            return []
        typ, idx = ph.get("type", "body"), ph.get("idx")
        sources = []
        layout = slide.slide_layout
        for holder, by_idx in ((layout, True), (layout.slide_master, False)):
            match = None
            for sh in holder.placeholders:
                sph = sh._element.find(".//" + qn("p:ph"))
                if sph is None:
                    continue
                if by_idx and idx is not None and sph.get("idx") == idx:
                    match = sh
                    break
                stype = sph.get("type", "body")
                same = {"ctrTitle": "title", "subTitle": "body", "obj": "body"}
                if same.get(stype, stype) == same.get(typ, typ) and match is None:
                    match = sh
            if match is not None:
                sources.append(match._element)
        return sources

    def placeholder_chain(self, el, slide):
        """([lstStyle containers of the layout and master placeholders], kind, [their bodyPr])."""
        qn = self.qn
        ph = el.find("./*/" + qn("p:nvPr") + "/" + qn("p:ph"))
        if ph is None:
            return [], "other", []
        kind = "title" if ph.get("type", "body") in ("title", "ctrTitle") else "body"
        containers, bodies = [], []
        for src in self.placeholder_sources(el, slide):
            tb = src.find(qn("p:txBody"))
            if tb is not None:
                containers.append(tb.find(qn("a:lstStyle")))
                bodies.append(tb.find(qn("a:bodyPr")))
        return containers, kind, bodies

    def shape_xfrm(self, el, slide):
        """A shape's own <a:xfrm>, or, for a placeholder without one, the nearest layout/master position."""
        qn = self.qn
        path = qn("p:spPr") + "/" + qn("a:xfrm")
        xfrm = el.find(path)
        if xfrm is None:
            for src in self.placeholder_sources(el, slide):
                xfrm = src.find(path)
                if xfrm is not None:
                    break
        return xfrm

    @staticmethod
    def _lvl(container, level):
        if container is None:
            return None
        return container.find(container.tag.rsplit("}", 1)[0] + "}" + f"lvl{level + 1}pPr")

    @staticmethod
    def _attr(chain, attr):
        for e in chain:
            if e is not None and e.get(attr) is not None:
                return e.get(attr)
        return None

    @staticmethod
    def _child(chain, *tags):
        for e in chain:
            if e is None:
                continue
            for c in e:
                if _local(c) in tags:
                    return _local(c), c
        return None, None

    @staticmethod
    def _spacing(el):
        if el is None:
            return None
        for c in el:
            t = _local(c)
            if t == "spcPct":
                return "pct", _int(c.get("val"), 100000) / 100000
            if t == "spcPts":
                return "pts", _int(c.get("val")) / 100
        return None

    def emit_text(self, el, txBody, box: _Box, style, ctx: dict):
        qn = self.qn
        theme: _Theme = ctx["theme"]
        bodyPr = txBody.find(qn("a:bodyPr"))
        containers, kind, ph_bodies = self.placeholder_chain(el, ctx["slide"])
        body_chain = [bodyPr] + ph_bodies

        def inset(attr, default):
            v = self._attr(body_chain, attr)
            return _int(v, default) if v is not None else default

        l_in, r_in = inset("lIns", 91440), inset("rIns", 91440)
        t_in, b_in = inset("tIns", 45720), inset("bIns", 45720)
        anchor = self._attr(body_chain, "anchor") or "t"
        wrap_text = (self._attr(body_chain, "wrap") or "square") != "none"
        vert = self._attr(body_chain, "vert") or "horz"
        font_scale, spacing_scale = 1.0, 1.0
        autofit = bodyPr.find(qn("a:normAutofit")) if bodyPr is not None else None
        if autofit is not None:
            font_scale = _int(autofit.get("fontScale"), 100000) / 100000
            spacing_scale = 1 - _int(autofit.get("lnSpcReduction"), 0) / 100000

        master_style = {"title": theme.title_style, "body": theme.body_style}.get(kind, theme.other_style)
        lst = txBody.find(qn("a:lstStyle"))
        font_ref = style.find(qn("a:fontRef")) if style is not None else None
        ref_color = theme.first_clr(font_ref) if font_ref is not None else None
        ref_font = theme.fonts.get(font_ref.get("idx")) if font_ref is not None and font_ref.get("idx") in ("major", "minor") else None
        default_text = ctx["default_text"]
        NB = "\xa0"
        # a link on the whole shape (Insert ▸ Link with the box selected) applies to text without a link of its own
        cnv = el.find(".//" + qn("p:cNvPr"))
        shape_link = self.link_url(cnv.find(qn("a:hlinkClick")) if cnv is not None else None, ctx)

        # the box as laid out: vertical text is laid out in a box turned on its side, then rotated back
        angle = box.angle
        bw, bh = box.w, box.h
        if vert in ("vert", "eaVert"):
            angle += 90
            bw, bh = bh, bw
        elif vert in ("vert270", "wordArtVert"):
            angle -= 90
            bw, bh = bh, bw
        inner_w = max(1.0, bw - l_in - r_in)

        # ---- lay the paragraphs out: each wrapped to the box width, with its own line height and spacing
        laid = []
        counters: dict = {}
        for p_el in txBody.findall(qn("a:p")):
            pPr = p_el.find(qn("a:pPr"))
            level = min(_int(pPr.get("lvl")) if pPr is not None else 0, 8)
            own = [pPr, self._lvl(lst, level)]
            inh = [self._lvl(c, level) for c in containers] + [self._lvl(master_style, level)]
            chain = own + inh
            runs = self._runs(p_el)
            end_rpr = p_el.find(qn("a:endParaRPr"))

            def rprops(rPr):
                srcs_own = [rPr] + [c.find(qn("a:defRPr")) if c is not None else None for c in own]
                srcs_inh = [c.find(qn("a:defRPr")) if c is not None else None for c in inh]
                srcs = srcs_own + srcs_inh
                sz = self._attr(srcs, "sz")
                size = (_int(sz) / 100 if sz else 18.0) * font_scale
                bold = self._attr(srcs, "b") == "1"
                italic = self._attr(srcs, "i") == "1"
                _, fill_own = self._child(srcs_own, "solidFill")
                hl = rPr.find(qn("a:hlinkClick")) if rPr is not None else None
                link = hl is not None
                color = None
                if fill_own is not None:
                    c = theme.first_clr(fill_own)
                    color = c[0] if c else None
                if color is None and link and theme.colors.get("hlink"):
                    color = "#" + theme.colors["hlink"]
                if color is None and ref_color:
                    color = ref_color[0]
                if color is None:
                    _, fill_inh = self._child(srcs_inh, "solidFill")
                    c = theme.first_clr(fill_inh) if fill_inh is not None else None
                    color = c[0] if c else None
                _, lat_own = self._child(srcs_own, "latin")
                font = theme.font(lat_own.get("typeface")) if lat_own is not None else None
                if not font and ref_font:
                    font = ref_font
                if not font:
                    _, lat_inh = self._child(srcs_inh, "latin")
                    font = theme.font(lat_inh.get("typeface")) if lat_inh is not None else None
                u = self._attr(srcs, "u")
                underline = (u is None and link) or (u is not None and u != "none")
                strike = self._attr(srcs, "strike") not in (None, "noStrike")
                return {"size": size, "bold": bold, "italic": italic, "color": color or default_text,
                        "font": font or theme.fonts["minor"], "underline": underline, "strike": strike,
                        "spc": _int(self._attr(srcs_own, "spc")), "url": self.link_url(hl, ctx)}

            main_rpr = next((rp for t, rp, _ in runs if t.strip()), None)
            mp = rprops(main_rpr if main_rpr is not None else end_rpr)
            size_pt, bold, italic, color, font = mp["size"], mp["bold"], mp["italic"], mp["color"], mp["font"]
            fs = max(6.0, size_pt * self.s_ref / PT_PER_FONT_SIZE)
            self.fonts_used.add(font)
            algn = self._attr(chain, "algn") or "l"
            mar_l = _int(self._attr(chain, "marL"))
            indent = _int(self._attr(chain, "indent"))
            spc_bef = self._spacing(self._child(chain, "spcBef")[1])
            spc_aft = self._spacing(self._child(chain, "spcAft")[1])
            ln_spc = self._spacing(self._child(chain, "lnSpc")[1])
            math_els = [m for _, _, m in runs if m is not None]
            math_only = bool(math_els) and not "".join(t for t, _, m in runs if m is None).strip()

            def measure(seg, st, font=font):
                ft = _font_or_fallback(font, *st)
                return ft.getlength(seg) / 1000 if ft is not None else len(seg) * (0.55 if st[0] else 0.5)

            # per-character style, colour and decoration; characters without a run take the paragraph's
            raw, styles, colors, decos, urls = "", [], [], [], []
            for t, rp, m in runs:
                t = t.replace("\x0b", "\n").replace("\t", "    ")
                st = col = deco = url = None
                spc = 0
                if rp is not None and t.strip():
                    rpp = rprops(rp)
                    st, col, deco, spc = (rpp["bold"], rpp["italic"]), rpp["color"], (rpp["underline"], rpp["strike"]), rpp["spc"]
                    url = rpp["url"]
                    if m is not None:
                        st = (rpp["bold"], rpp["italic"] and m.find(".//" + _MATH_NS + "nor") is None)
                if spc >= 50 and t.strip():
                    # letter spacing: manim's Text has none, so spread the characters with no-break spaces
                    sp_pt = measure(" ", st or (bold, italic)) * size_pt
                    n = round(spc / 100 / max(sp_pt, 0.1))
                    if n >= 1:
                        t = (NB * n).join(t)
                raw += t
                styles += [st] * len(t)
                colors += [col] * len(t)
                decos += [deco] * len(t)
                urls += [url or shape_link] * len(t)
            lead = len(raw) - len(raw.lstrip())
            text = raw.strip()
            styles, colors, decos = styles[lead:lead + len(text)], colors[lead:lead + len(text)], decos[lead:lead + len(text)]
            urls = urls[lead:lead + len(text)]
            counts: dict = {}
            for ch, st in zip(text, styles):
                if st is not None and not ch.isspace():
                    counts[st] = counts.get(st, 0) + 1
            if counts:
                bold, italic = max(counts, key=counts.get)
            styles = [st or (bold, italic) for st in styles]
            colors = [c or color for c in colors]
            decos = [d or (mp["underline"], mp["strike"]) for d in decos]
            empty = text == "" and not math_only

            # bullets: drawn as their own small Text in front of the paragraph's first line
            bullet, bullet_color, bullet_scale, bullet_font = "", color, 1.0, font
            if not empty:
                btag, bel = self._child(chain, "buNone", "buChar", "buAutoNum", "buBlip")
                if btag == "buChar":
                    bfont_el = self._child(chain, "buFont")[1]
                    bfont = bfont_el.get("typeface") if bfont_el is not None else ""
                    bullet = _BULLETS.get((bfont, bel.get("char", "•")), bel.get("char", "•"))
                    if bfont and (bfont, bel.get("char", "•")) not in _BULLETS and not bfont.startswith(("Wingdings", "Webdings", "Symbol")):
                        bullet_font = theme.font(bfont) or font
                elif btag == "buAutoNum":
                    scheme = bel.get("type", "arabicPeriod")
                    key = (level, scheme)
                    n = counters.get(key, _int(bel.get("startAt"), 1) - 1) + 1
                    counters[key] = n
                    if scheme.startswith("alphaLc"):
                        s = chr(ord("a") + (n - 1) % 26)
                    elif scheme.startswith("alphaUc"):
                        s = chr(ord("A") + (n - 1) % 26)
                    elif scheme.startswith("roman"):
                        s = self._roman(n)
                        s = s.lower() if scheme.startswith("romanLc") else s
                    else:
                        s = str(n)
                    bullet = s + (")" if scheme.endswith("ParenR") or scheme.endswith("ParenBoth") else "." if scheme.endswith("Period") else "")
                    if scheme.endswith("ParenBoth"):
                        bullet = "(" + bullet
                elif btag == "buNone":
                    counters = {k: v for k, v in counters.items() if k[0] != level}
                if bullet:
                    bc = self._child(chain, "buClr")[1]
                    c = theme.first_clr(bc) if bc is not None else None
                    if c:
                        bullet_color = c[0]
                    bsz = self._child(chain, "buSzPct")[1]
                    if bsz is not None:
                        bullet_scale = _int(bsz.get("val"), 100000) / 100000
            bullet_w = measure(bullet, (bold, italic)) * size_pt * 12700 * bullet_scale if bullet else 0.0
            if bullet:
                bullet_x = mar_l + indent
                text_x = mar_l if indent < 0 else bullet_x + bullet_w + 91440
            else:
                bullet_x, text_x = 0.0, mar_l + indent
            text_x = max(0.0, text_x)

            avail_em = max(0.5, (inner_w - text_x) / 12700 / size_pt)
            if wrap_text and not empty and not math_only:
                text, styles, idx = _wrap(text, styles, avail_em, measure)
                colors = [colors[i] if i >= 0 else color for i in idx]
                decos = [decos[i] if i >= 0 else (False, False) for i in idx]
                urls = [urls[i] if i >= 0 else None for i in idx]
            text_lines = text.split("\n")

            pil = _font_or_fallback(font, bold, italic)
            if pil is not None:
                ascent, descent = pil.getmetrics()
                line_em = getattr(pil.font, "height", ascent + descent) / 1000
                asc_em, desc_em = ascent / 1000, descent / 1000
            else:
                line_em, asc_em, desc_em = 1.2, 0.95, 0.25
            size_emu = size_pt * 12700
            if ln_spc and ln_spc[0] == "pts":
                pitch = ln_spc[1] * 12700 * spacing_scale
            else:
                pitch = (ln_spc[1] if ln_spc else 1.0) * spacing_scale * line_em * size_emu

            def reach(s, which):
                # how far the glyphs of a line rise above (which=0) / fall below (which=1) its baseline, in ems
                if pil is None or not s.strip():
                    return 0.7 if which == 0 else 0.0
                bb = pil.getbbox(s)
                return (ascent - bb[1]) / 1000 if which == 0 else (bb[3] - ascent) / 1000

            def sp(v):
                if v is None:
                    return 0.0
                return v[1] * 12700 if v[0] == "pts" else v[1] * line_em * size_emu

            widths, off = [], 0
            for ln in text_lines:
                widths.append(self._seg_width(text, styles, off, off + len(ln), measure))
                off += len(ln) + 1
            latex = ""
            n_lines = float(len(text_lines))
            if math_only:
                try:
                    latex = " ".join(omml_latex(m) for m in math_els) or "?"
                except Exception:                               # noqa: BLE001
                    latex = ""
                n_lines = 2.2 if any(omml_tall(m) for m in math_els) else 1.0
            laid.append({"text": text, "styles": styles, "colors": colors, "decos": decos, "urls": urls, "n_lines": n_lines, "fs": fs,
                         "size_emu": size_emu, "color": color, "font": font, "pitch": pitch, "align": algn,
                         "empty": empty, "text_x": text_x, "bold": bold, "italic": italic, "math": latex,
                         "bullet": bullet, "bullet_x": bullet_x, "bullet_w": bullet_w, "bullet_color": bullet_color, "bullet_font": bullet_font,
                         "bullet_scale": bullet_scale, "before": sp(spc_bef), "after": sp(spc_aft),
                         "desc_em": desc_em, "top_reach": reach(text_lines[0], 0),
                         "bottom_reach": reach(text_lines[-1], 1), "widths": widths, "measure": measure,
                         "pil": pil, "ascent": pil.getmetrics()[0] if pil is not None else 950})
        if not laid:
            return
        for p in laid:
            linked = [u for ch, u in zip(p["text"], p["urls"]) if u and not ch.isspace()]
            p["link"] = max(set(linked), key=linked.count) if linked else None

        # ---- group paragraphs into Texts: one per text box when the paragraphs share size, font and alignment
        def left_aligned(p):
            return p["align"] in ("l", "just", "dist")

        def compatible(g, p):
            ref = next((q for q in g if not q["empty"]), None)
            if ref is None or p["empty"]:
                return True
            if p["math"] or ref["math"]:
                return False
            # a link belongs to a whole Text, so paragraphs that open different links stay apart
            if any(q["link"] and p["link"] and q["link"] != p["link"] for q in g):
                return False
            return (p["fs"] == ref["fs"] and p["font"] == ref["font"] and p["align"] == ref["align"]
                    and (not left_aligned(p) or abs(p["text_x"] - ref["text_x"]) < 1000))

        groups: list = []
        for p in laid:
            if groups and compatible(groups[-1], p):
                groups[-1].append(p)
            else:
                groups.append([p])

        def group_height(g):
            h = sum(p["n_lines"] * p["pitch"] for p in g)
            h += sum(p["after"] for p in g[:-1]) + sum(p["before"] for p in g[1:])
            return h

        total = sum(group_height(g) for g in groups)
        total += sum(g[0]["before"] for g in groups[1:]) + sum(g[-1]["after"] for g in groups[:-1])
        inner_h = bh - t_in - b_in
        if anchor == "t":
            y = 0.0
        elif anchor == "b":
            y = inner_h - total
        else:
            y = (inner_h - total) / 2
        y += t_in                                                 # distance from the top edge of the box, downwards

        rot_rad = -math.radians(angle)
        a_rad = math.radians(angle)
        out = ctx["lines"]

        def at(dx, dy):
            """Slide position (EMU) of a point given relative to the box centre in the laid-out frame."""
            if abs(angle) < 0.05:
                return box.cx + dx, box.cy + dy
            return (box.cx + dx * math.cos(a_rad) - dy * math.sin(a_rad),
                    box.cy + dx * math.sin(a_rad) + dy * math.cos(a_rad))

        def place(dx, dy, left_x=None, right_x=None):
            """Positioning code for a mobject whose centre is at (dx, dy) in the laid-out frame; when the box is
            not rotated, its left/right edge is pinned exactly instead of relying on the estimated width."""
            x, yy = at(dx, dy)
            pos = f".move_to(({_num(self.ux(x))}, {_num(self.uy(yy))}, 0))"
            if abs(angle) < 0.05:
                if left_x is not None:
                    pos += f".align_to(({_num(self.ux(box.cx - bw / 2 + left_x))}, 0, 0), LEFT)"
                elif right_x is not None:
                    pos += f".align_to(({_num(self.ux(box.cx - bw / 2 + right_x))}, 0, 0), RIGHT)"
            else:
                pos += f".rotate({_num(rot_rad, 4)})"
            return pos

        def emit_bullet(p, baseline, fs_text):
            sz = p["size_emu"] * p["bullet_scale"]
            pil = p["pil"]
            if pil is not None and p["bullet"].strip():
                bb = pil.getbbox(p["bullet"])
                top_r, bot_r = (p["ascent"] - bb[1]) / 1000, (bb[3] - p["ascent"]) / 1000
            else:
                top_r, bot_r = 0.6, 0.0
            cy_emu = baseline + (bot_r - top_r) / 2 * sz
            dx = -bw / 2 + l_in + p["bullet_x"] + p["bullet_w"] / 2
            dy = cy_emu - bh / 2
            out.append(f'self.add(Text({p["bullet"]!r}, font_size={_num(fs_text * p["bullet_scale"], 1)}, color="{p["bullet_color"]}", '
                       f'font="{p["bullet_font"]}", disable_ligatures=True){place(dx, dy, left_x=l_in + p["bullet_x"])})  # manimble')

        for gi, g in enumerate(groups):
            if gi:
                y += g[0]["before"]
            block_h = group_height(g)
            first = next((i for i, p in enumerate(g) if not p["empty"]), None)
            if first is None:
                y += block_h + g[-1]["after"]
                continue
            last = max(i for i, p in enumerate(g) if not p["empty"])
            # the Text leaves out empty paragraphs at either end, but they take up space
            lead_h = sum(p["n_lines"] * p["pitch"] + p["after"] for p in g[:first]) + sum(p["before"] for p in g[1:first + 1])
            mid = g[first:last + 1]
            mid_h = group_height(mid)
            ref = mid[0]
            size_emu = ref["size_emu"]
            fs = ref["fs"]
            align = ref["align"]

            if ref["math"]:
                # a formula of its own: MathTex, left-aligned like the text would be (its size is scaled so that
                # LaTeX's letters come out about as tall as the text font's)
                yc = y + lead_h + mid_h / 2
                dy = yc - bh / 2
                if left_aligned(ref):
                    dx = -bw / 2 + l_in + ref["text_x"] + 0.3 * size_emu
                    pos = place(dx, dy, left_x=l_in + ref["text_x"])
                elif align == "r":
                    pos = place(bw / 2 - r_in - 0.3 * size_emu, dy, right_x=bw - r_in)
                else:
                    pos = place((l_in - r_in) / 2, dy)
                if ref["bullet"]:
                    emit_bullet(ref, yc, fs)
                url = next((u for u in ref["urls"] if u), shape_link)
                pos += f".set(manimble_link={url!r})" if url else ""
                out.append(f'self.add(MathTex({ref["math"]!r}, font_size={_num(fs * 1.3, 1)}, color="{ref["color"]}"){pos})  # manimble')
                y += block_h + g[-1]["after"]
                continue

            n_lines = int(sum(p["n_lines"] for p in mid))
            pitch_eff = mid_h / n_lines if n_lines else ref["pitch"]
            # glyph top of the first line and glyph bottom of the last line, from the block's top (EMU, down)
            first_baseline = y + lead_h + ref["pitch"] - ref["desc_em"] * size_emu
            glyph_top = first_baseline - ref["top_reach"] * size_emu
            glyph_bottom = first_baseline + (n_lines - 1) * pitch_eff + mid[-1]["bottom_reach"] * size_emu
            yc_emu = (glyph_top + glyph_bottom) / 2
            y += block_h + g[-1]["after"]

            text = "\n".join(p["text"] for p in mid)
            styles, colors, decos, urls = [], [], [], []
            for pi, p in enumerate(mid):
                more = pi < len(mid) - 1
                urls += p["urls"]
                styles += p["styles"] + ([(p["bold"], p["italic"])] if more else [])
                colors += p["colors"] + ([p["color"]] if more else [])
                decos += p["decos"] + ([(False, False)] if more else [])

            def main(values):
                counts: dict = {}
                for ch, v in zip(text, values):
                    if not ch.isspace():
                        counts[v] = counts.get(v, 0) + 1
                return max(counts, key=counts.get)

            def ranges(values, base):
                # manim's "[i:j]" ranges index the text as written, spaces and line breaks included
                rs, i = [], 0
                while i < len(text):
                    j = i
                    while j < len(text) and values[j] == values[i]:
                        j += 1
                    if values[i] != base and text[i:j].strip():
                        rs.append((i, j, values[i]))
                    i = j
                return rs

            color = main(colors)
            bold = main([st[0] for st in styles])
            italic = main([st[1] for st in styles])
            kw = [f"font_size={_num(fs, 1)}", f'color="{color}"', f'font="{ref["font"]}"']
            if bold:
                kw.append('weight="BOLD"')
            if italic:
                kw.append('slant="ITALIC"')
            t2c = ranges(colors, color)
            t2w = [(i, j, "BOLD" if v else "NORMAL") for i, j, v in ranges([st[0] for st in styles], bold)]
            t2s = [(i, j, "ITALIC" if v else "NORMAL") for i, j, v in ranges([st[1] for st in styles], italic)]
            for key, rs in (("t2c", t2c), ("t2w", t2w), ("t2s", t2s)):
                if rs:
                    kw.append(f"{key}={{" + ", ".join(f'"[{i}:{j}]": "{v}"' for i, j, v in rs) + "}")
            total_lines = text.count("\n") + 1
            if total_lines > 1:
                # manim's line pitch is font_size * (1 + line_spacing) / 96 units
                kw.append(f"line_spacing={_num(self.uh(pitch_eff) * 96 / fs - 1, 2)}")
            kw.append("disable_ligatures=True")                 # keeps t2c/t2w ranges and Paragraph lines aligned
            if total_lines > 1 and align in ("ctr", "r"):
                lines_lit = ", ".join(repr(l) for l in text.split("\n"))
                call = f'Paragraph({lines_lit}, alignment="{"center" if align == "ctr" else "right"}", {", ".join(kw)})'
            else:
                call = f"Text({text!r}, {', '.join(kw)})"

            # position: offsets from the box centre in the laid-out frame (EMU, x right, y down)
            widths = [w for p in mid for w in p["widths"]]
            widest = max(widths) * size_emu if widths else 0.0
            dy = yc_emu - bh / 2
            if left_aligned(ref):
                pos = place(-bw / 2 + l_in + ref["text_x"] + widest / 2, dy, left_x=l_in + ref["text_x"])
            elif align == "r":
                pos = place(bw / 2 - r_in - widest / 2, dy, right_x=bw - r_in)
            else:
                pos = place((l_in - r_in) / 2, dy)
            # a link belongs to the whole Text in Manimble: the one on most of its linked characters
            linked: dict = {}
            for ch, u in zip("".join(p["text"] for p in mid), urls):
                if u and not ch.isspace():
                    linked[u] = linked.get(u, 0) + 1
            if linked:
                url = max(linked, key=linked.get)
                pos += f".set(manimble_link={url!r})"
                if len(linked) > 1:
                    dropped = ", ".join(u for u in linked if u != url)
                    self.warnings.append(f"slide {ctx['slide_idx']}: a line with several links keeps only {url} "
                                         f"(dropped {dropped})")
            decorated = any(d[0] or d[1] for d in decos) and abs(angle) < 0.05
            var = None
            if decorated:
                # underlines are anchored to the Text's own glyphs, so the Text gets a name
                ctx["n_text"] = ctx.get("n_text", 0) + 1
                var = f"t{ctx['n_text']}"
                out.append(f"{var} = {call}{pos}  # manimble")
                out.append(f"self.add({var})  # manimble")
            else:
                out.append(f"self.add({call}{pos})  # manimble")
            is_paragraph = call.startswith("Paragraph(")

            # bullets, underlines and strikethroughs follow the lines of the Text
            line_no = 0
            for p in mid:
                if p["bullet"]:
                    emit_bullet(p, first_baseline + line_no * pitch_eff, fs)
                line_no += int(p["n_lines"])
            if not decorated:
                continue
            off = 0
            for k, ln in enumerate(text.split("\n")):
                baseline = first_baseline + k * pitch_eff
                i = 0
                while i < len(ln):
                    d = decos[off + i]
                    if not (d[0] or d[1]) or ln[i].isspace():
                        i += 1
                        continue
                    j = i
                    while j < len(ln) and decos[off + j] == d:
                        j += 1
                    while j > i and ln[j - 1].isspace():
                        j -= 1
                    # x from the glyphs of the rendered Text: with disable_ligatures, .chars has one entry per
                    # character of the text as written (spaces and line breaks included, as empty glyphs)
                    def glyph(pos_in_line):
                        return f"{var}.chars[{k}][{pos_in_line}]" if is_paragraph else f"{var}.chars[{off + pos_in_line}]"
                    xa = f"{glyph(i)}.get_left()[0]"
                    xb = f"{glyph(j - 1)}.get_right()[0]"
                    sw = _num(max(1.0, 0.07 * size_emu * self.px_per_emu / 0.9), 1)
                    for on, dyl in ((d[0], 0.11 * size_emu), (d[1], -0.28 * size_emu)):
                        if not on:
                            continue
                        ya = _num(self.uy(box.cy + baseline + dyl - bh / 2))
                        out.append(f'self.add(Line(({xa}, {ya}, 0), ({xb}, {ya}, 0), color="{colors[off + i]}", stroke_width={sw}))  # manimble')
                    i = j
                off += len(ln) + 1

    def link_url(self, hl, ctx: dict) -> str | None:
        """The web/mail address an a:hlinkClick opens; None for jumps to slides and other actions, which a show in
        Manimble cannot follow."""
        if hl is None or (hl.get("action") or "").startswith("ppaction://"):
            return None
        rid = hl.get(self.qn("r:id"))
        try:
            rel = ctx["slide"].part.rels[rid]
        except (KeyError, TypeError):
            return None
        return rel.target_ref if rel.is_external and rel.target_ref else None

    @staticmethod
    def _seg_width(text, styles, start, end, measure):
        """Width of text[start:end] in ems; styles are indexed like `text`."""
        w, k = 0.0, start
        while k < end:
            m = k
            while m < end and styles[m] == styles[k]:
                m += 1
            w += measure(text[k:m], styles[k])
            k = m
        return w

    @staticmethod
    def _roman(n: int) -> str:
        out = ""
        for v, s in ((1000, "M"), (900, "CM"), (500, "D"), (400, "CD"), (100, "C"), (90, "XC"), (50, "L"), (40, "XL"),
                     (10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")):
            while n >= v:
                out += s
                n -= v
        return out

    # ----------------------------------------------------------------------------------------------- slides
    def run(self, src_name: str, progress=None) -> dict:
        qn = self.qn
        self.deck.mkdir(parents=True, exist_ok=True)
        self.assets.mkdir(parents=True, exist_ok=True)
        if abs(self.W / self.H - 16 / 9) > 0.02:
            self.warnings.append("the presentation is not 16:9; it was stretched onto a 16:9 frame")
        scene_names: list[str] = []
        classes: list[str] = []
        slides = list(self.prs.slides)
        for si, slide in enumerate(slides, start=1):
            if progress:
                progress(si - 1, len(slides), f"slide {si}")
            theme = self.theme(slide)
            bg = self.slide_bg(slide, theme)
            dark = _luma(bg) < 128
            text_default = theme.colors.get(theme.clr_map.get("tx1", "dk1"))
            ctx = {"slide_idx": si, "slide": slide, "theme": theme, "bg": bg, "lines": [], "group_fill": None,
                   "default_text": ("#" + text_default) if text_default else ("#FFFFFF" if dark else "#000000")}
            title = ""
            for shape in self.iter_shapes(slide.shapes, slide.shapes._spTree):
                if not title and shape.has_text_frame and shape.text_frame.text.strip():
                    title = shape.text_frame.text.strip().split("\n")[0]
                try:
                    self.emit_shape(shape, _Xf(), ctx)
                except Exception as exc:                        # noqa: BLE001
                    tb = exc.__traceback__
                    while tb is not None and tb.tb_next is not None:
                        tb = tb.tb_next
                    where = f" at importer.py:{tb.tb_lineno}" if tb is not None else ""
                    self.warnings.append(f"slide {si}: skipped a shape ({exc}{where})")
            name = f"S{si:02d}_{_slug(title, 'Slide')}"
            scene_names.append(name)
            notes = ""
            try:
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                    notes = "\n        # Speaker notes:\n" + "\n".join("        # " + l for l in slide.notes_slide.notes_text_frame.text.strip().split("\n"))
            except Exception:                                   # noqa: BLE001
                pass
            code = [f"class {name}(Deck):", f'    """Imported from slide {si}."""', "", "    def construct(self):",
                    f'        self.camera.background_color = "{bg}"{notes}']
            code += ["        " + l for l in ctx["lines"]] or ["        pass"]
            classes.append("\n".join(code))

        fonts = sorted(f for f in self.fonts_used if f)
        register = ""
        if fonts and OFFICE_FONT_DIRS:
            register = FONT_REGISTER_SOURCE + "\n\n_register_fonts(" + ", ".join(repr(f) for f in fonts) + ")\n"
        (self.deck / "common.py").write_text(
            '"""Shared base class for the imported deck."""\n\nfrom pathlib import Path\n\nfrom manim import *\nfrom manim_slides import Slide\n\n'
            'ASSETS = Path(__file__).resolve().parent.parent / "assets"\n\n\nclass Deck(Slide):\n    skip_reversing = True\n'
            + (TABLE_GRID_SOURCE if self.uses_tables else "") + register)
        (self.deck / "slides.py").write_text(
            '"""Slides imported from ' + src_name + '."""\n\nfrom common import *\n\n\n' + "\n\n\n".join(classes) + "\n")
        (self.deck / "order.txt").write_text(" ".join(scene_names) + "\n")
        (self.deck / "manim.cfg").write_text("[CLI]\nbackground_color = #000000\npixel_width = 1280\npixel_height = "
                                             + str(int(round(1280 * self.H / self.W))) + "\nframe_rate = 30\n")
        seen: set = set()
        warnings = [w for w in self.warnings if not (w in seen or seen.add(w))]
        return {"ok": True, "deck": str(self.deck), "slides": len(scene_names), "warnings": warnings[:20]}


def import_presentation(src: str, dest: str, progress=None) -> dict:
    """Create dest/deck (+ dest/assets) from the presentation. Returns {"ok", "deck", "slides", "warnings"}."""
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.oxml.ns import qn
    except ImportError:
        return {"ok": False, "error": "python-pptx is not installed in the Python environment Manimble uses "
                                      "(pip install python-pptx)."}
    src_p = Path(src)
    if src_p.suffix.lower() in (".key", ".ppt"):         # python-pptx reads only .pptx; Keynote converts
        tmp = Path(tempfile.mkdtemp()) / (src_p.stem + ".pptx")
        r = keynote_to_pptx(str(src_p), str(tmp))
        if not r.get("ok"):
            return r
        src_p = tmp
    try:
        prs = Presentation(str(src_p))
    except Exception as exc:                                    # noqa: BLE001
        return {"ok": False, "error": f"Could not read the presentation: {exc}"}
    return _Importer(prs, Path(dest), qn, (MSO_SHAPE_TYPE, PP_ALIGN, MSO_ANCHOR)).run(src_p.name, progress)
